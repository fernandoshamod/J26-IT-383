"""Versioned observation-only refinement, independent of synthetic approval gates.

No imputation, source substitution, inflation/price merging or fitted transformations.
"""
from collections import defaultdict
import csv
from datetime import date
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import subprocess
import tempfile

from .historical_audit import audit_project, read_csv, sha, month_number

VERSION = 'historical-refinement.0.1'
OBSERVATION_FIELDS = ['observation_id', 'ds', 'y', 'series_id', 'source_id', 'category',
                      'unit', 'observation_kind', 'synthetic_origin', 'geographic_context',
                      'geographic_context_evidence', 'publication_date', 'lineage_reconciled',
                      'closed_calendar_month', 'source_quality_review', 'quality_flag',
                      'source_record_count', 'reported_weeks', 'snapshot_available_weeks',
                      'covers_available_snapshot_weeks', 'publisher_coverage_status',
                      'technical_split_screen_pass', 'model_use_status']
LINEAGE_FIELDS = ['observation_id', 'source_record_id', 'source_id', 'source_extract_file',
                  'source_csv_row', 'source_extract_sha256', 'source_quality_flag',
                  'source_observed_value', 'source_unit', 'reported_week', 'publisher_filename',
                  'publisher_url', 'publisher_sha256', 'publisher_page', 'publisher_bytes_verified']
SPLIT_FIELDS = ['observation_id', 'series_id', 'ds', 'partition', 'split_status',
                'source_quality_review', 'model_training_authorized']
SOURCE_PATHS = ['data/processed/inflation.csv', 'data/processed/retail_prices.csv',
                'data/final/prophet_cost_timeseries.csv', 'data/raw/extractions/Inflation.csv',
                'data/raw/extractions/Retail_Prices.csv',
                'outputs/01a12103-7a09-7331-bb93-9e24a6f05005/download_manifest.json',
                'work/original_schema.json']


def _positive(value):
    try:
        return math.isfinite(float(value)) and float(value) > 0
    except (ValueError, TypeError):
        return False


def expected_observations(inflation, retail):
    """Independently reconstruct the complete legacy extraction selection recipe."""
    expected, groups = {}, defaultdict(list)
    for r in inflation:
        if not r['index_value']:
            continue
        key = (f"{r['source_id']}:{r['index_type']}:base{r['base_year']}", r['month'])
        if key in expected:
            raise ValueError('Conflicting index source month')
        expected[key] = {r['source_record_id']}
    for r in retail:
        if (r['parse_status'] != 'Parsed' or not _positive(r['current_price_lkr']) or
                not r['unit_raw'] or re.match(r'^(Imported|Local|No[ .]|Red$|White$)', r['item_raw'], re.I)):
            continue
        groups[(r['year'], r['month'], r['item_raw'], r['unit_raw'])].append(r)
    for (year, month, item, unit), records in groups.items():
        if len({r['reported_week'] for r in records}) != len(records):
            continue  # Ambiguous PDF grouping remains excluded, not arbitrarily deduplicated.
        key = (f'S003:{item}:{unit}', f'{int(year):04d}-{int(month):02d}-01')
        if key in expected:
            raise ValueError('Conflicting retail series identity')
        expected[key] = {r['source_record_id'] for r in records}
    return expected


def _publisher_map(root, download_manifest):
    mapping, file_hashes = {}, {}
    for row in download_manifest:
        if row['source_id'] not in {'S001', 'S002', 'S003'} or row['status'] != 'Downloaded':
            continue
        name = row['name']
        if Path(name).name != name:
            raise ValueError('Unsafe publisher filename')
        path = root / 'data/raw/publishers' / row['source_id'] / name
        key = (row['source_id'], name)
        if key in mapping:
            raise ValueError('Ambiguous publisher manifest identity')
        digest = sha(path)
        if digest != row['sha256'] or path.stat().st_size != row['bytes']:
            raise ValueError('Publisher snapshot checksum or length mismatch')
        mapping[key] = row
        file_hashes[str(path.relative_to(root)).replace('\\', '/')] = digest
    return mapping, file_hashes


def build(root, as_of):
    root = Path(root).resolve()
    date.fromisoformat(as_of)
    audit = audit_project(root, as_of)
    if (audit['reconciled_rows'] != audit['rows'] or audit['extraction_identity_errors'] or
            audit['extraction_content_errors'] or not audit['original_inputs_preserved']):
        raise ValueError('Historical source audit failed; no refinement is permitted')
    inflation, retail = (read_csv(root / p) for p in SOURCE_PATHS[:2])
    existing = read_csv(root / SOURCE_PATHS[2])
    if not existing:
        raise ValueError('Empty historical observations are not a completed dataset')
    expected = expected_observations(inflation, retail)
    actual = {(r['series_id'], r['ds']): set(r['provenance'].split(';')) for r in existing}
    if actual != expected:
        raise ValueError('Export does not match the complete source selection and lineage recipe')
    manifest = json.loads((root / SOURCE_PATHS[5]).read_text(encoding='utf-8'))
    publishers, publisher_hashes = _publisher_map(root, manifest)
    sources = {r['source_record_id']: r for r in inflation + retail}
    summaries = {r['series_id']: r for r in audit['series']}
    available_weeks = defaultdict(set)
    for r in retail:
        available_weeks[(int(r['year']), int(r['month']))].add(r['reported_week'])
    cutoff = month_number(date.fromisoformat(as_of).replace(day=1).isoformat())
    observations, lineage = [], []
    for r in sorted(existing, key=lambda x: (x['series_id'], x['ds'])):
        refs = sorted(r['provenance'].split(';'))
        contributors = [sources[rid] for rid in refs]
        identity = json.dumps([r['series_id'], r['ds'], refs], separators=(',', ':'), ensure_ascii=False)
        oid = 'HIST-' + hashlib.sha256(identity.encode()).hexdigest()[:24]
        quality = sorted({flag for c in contributors for flag in c['quality_flag'].split(';') if flag})
        index = r['source_id'] in {'S001', 'S002'}
        closed = month_number(r['ds']) < cutoff
        weeks = None if index else len({c['reported_week'] for c in contributors})
        available = None if index else len(available_weeks[(int(r['ds'][:4]), int(r['ds'][5:7]))])
        # Geographic descriptions are explicitly sourced to the retained planning
        # register, not asserted as freshly verified publisher geographic definitions.
        geo = {'S001': 'Colombo', 'S002': 'Sri Lanka (national)', 'S003': 'Colombo District'}[r['source_id']]
        observations.append({'observation_id': oid, 'ds': r['ds'], 'y': r['y'],
                             'series_id': r['series_id'], 'source_id': r['source_id'],
                             'category': r['cost_category'], 'unit': r['unit'],
                             'observation_kind': r['observation_kind'], 'synthetic_origin': False,
                             'geographic_context': geo,
                             'geographic_context_evidence': 'original Data_Sources register; publisher scope review pending',
                             'publication_date': None, 'lineage_reconciled': True,
                             'closed_calendar_month': closed, 'source_quality_review': bool(quality),
                             'quality_flag': ';'.join(quality), 'source_record_count': len(refs),
                             'reported_weeks': weeks, 'snapshot_available_weeks': available,
                             'covers_available_snapshot_weeks': None if index else weeks == available,
                             'publisher_coverage_status': 'not independently established',
                             'technical_split_screen_pass': summaries[r['series_id']]['technical_split_screen_pass'],
                             'model_use_status': 'historical review candidate; no model fitted'})
        for rid, c in zip(refs, contributors):
            filename = ({'S001': 'CCPI_base2021.pdf', 'S002': 'NCPI_base2021.pdf'}[r['source_id']]
                        if index else c['source_file'])
            publisher = publishers.get((r['source_id'], filename))
            if publisher is None:
                raise ValueError(f'Missing publisher lineage: {r["source_id"]}/{filename}')
            lineage.append({'observation_id': oid, 'source_record_id': rid, 'source_id': r['source_id'],
                            'source_extract_file': f'data/raw/extractions/{"Inflation" if index else "Retail_Prices"}.csv',
                            'source_csv_row': int(c['source_csv_row']),
                            'source_extract_sha256': c['source_extract_sha256'],
                            'source_quality_flag': c['quality_flag'],
                            'source_observed_value': c['index_value'] if index else c['current_price_lkr'],
                            'source_unit': r['unit'], 'reported_week': None if index else c['reported_week'],
                            'publisher_filename': filename, 'publisher_url': publisher['url'],
                            'publisher_sha256': publisher['sha256'],
                            'publisher_page': None if index else int(c['pdf_page']),
                            'publisher_bytes_verified': True})
    partitions = []
    by_series = defaultdict(list)
    for r in observations:
        if r['closed_calendar_month'] and r['technical_split_screen_pass']:
            by_series[r['series_id']].append(r)
    for sid, rows in sorted(by_series.items()):
        ntrain, nvalidend = int(.7 * len(rows)), int(.85 * len(rows))
        for i, r in enumerate(rows):
            partitions.append({'observation_id': r['observation_id'], 'series_id': sid, 'ds': r['ds'],
                               'partition': 'train' if i < ntrain else 'validation' if i < nvalidend else 'test',
                               'split_status': 'proposed per-series retrospective univariate split',
                               'source_quality_review': r['source_quality_review'],
                               'model_training_authorized': False})
    result = {'version': VERSION, 'as_of': as_of, 'observations': observations,
              'closed_month_observations': [r for r in observations if r['closed_calendar_month']],
              'chronological_partitions_proposed': partitions, 'observation_provenance': lineage,
              'series_coverage': audit['series'], 'audit': audit,
              'source_hashes': {p: sha(root / p) for p in SOURCE_PATHS},
              'publisher_hashes': publisher_hashes,
              'model_training_authorized': False, 'synthetic_rows_generated': 0}
    validate(result)
    return result


def validate(bundle):
    observations = bundle['observations']
    ids, keys = set(), set()
    lookup = {}
    for r in observations:
        if set(r) != set(OBSERVATION_FIELDS):
            raise ValueError('Observation schema mismatch')
        key = (r['series_id'], r['ds'])
        if r['observation_id'] in ids or key in keys or not _positive(r['y']):
            raise ValueError('Duplicate identity/key or invalid measurement')
        month_number(r['ds'])
        if r['synthetic_origin'] is not False or r['lineage_reconciled'] is not True or r['publication_date'] is not None:
            raise ValueError('Invalid origin or invented publication date')
        ids.add(r['observation_id']); keys.add(key); lookup[r['observation_id']] = r
    if not observations or bundle['model_training_authorized'] is not False or bundle['synthetic_rows_generated'] != 0:
        raise ValueError('Invalid dataset or execution boundary')
    expected_closed = [r for r in observations if r['closed_calendar_month']]
    if bundle['closed_month_observations'] != expected_closed:
        raise ValueError('Closed-month export differs from observation filter')
    contributors = defaultdict(list)
    seen = set()
    for r in bundle['observation_provenance']:
        key = (r['observation_id'], r['source_record_id'])
        if set(r) != set(LINEAGE_FIELDS) or key in seen or r['observation_id'] not in ids or r['publisher_bytes_verified'] is not True:
            raise ValueError('Invalid provenance schema/identity/snapshot')
        seen.add(key); contributors[r['observation_id']].append(r)
    for oid, r in lookup.items():
        parts = contributors[oid]
        if len(parts) != r['source_record_count'] or any(p['source_id'] != r['source_id'] or p['source_unit'] != r['unit'] for p in parts):
            raise ValueError('Incomplete or mixed-unit/source provenance')
        values = [float(p['source_observed_value']) for p in parts]
        if not all(math.isfinite(v) and v > 0 for v in values) or not math.isclose(math.fsum(values) / len(values), float(r['y']), rel_tol=1e-12, abs_tol=1e-8):
            raise ValueError('Export does not reconcile to provenance')
    split_seen, dates = set(), defaultdict(lambda: defaultdict(list))
    for r in bundle['chronological_partitions_proposed']:
        if set(r) != set(SPLIT_FIELDS) or r['observation_id'] in split_seen or r['observation_id'] not in ids:
            raise ValueError('Duplicate or invalid split assignment')
        source = lookup[r['observation_id']]
        if (not source['closed_calendar_month'] or not source['technical_split_screen_pass'] or
                (r['series_id'], r['ds']) != (source['series_id'], source['ds']) or
                r['model_training_authorized'] is not False or r['partition'] not in {'train', 'validation', 'test'}):
            raise ValueError('Invalid chronological scope or authority')
        split_seen.add(r['observation_id']); dates[r['series_id']][r['partition']].append(r['ds'])
    expected_split = {r['observation_id'] for r in observations if r['closed_calendar_month'] and r['technical_split_screen_pass']}
    if split_seen != expected_split:
        raise ValueError('Missing eligible chronological assignment')
    for sid, parts in dates.items():
        if (set(parts) != {'train', 'validation', 'test'} or
                not max(parts['train']) < min(parts['validation']) or
                not max(parts['validation']) < min(parts['test'])):
            raise ValueError(f'Chronological leakage in {sid}')
    return True


def _csv(path, rows, fields):
    with path.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, lineterminator='\n')
        writer.writeheader()
        for row in rows:
            writer.writerow({k: json.dumps(row[k], ensure_ascii=False, separators=(',', ':'))
                             if isinstance(row[k], (list, dict)) else int(row[k])
                             if isinstance(row[k], bool) else row[k] for k in fields})


def verify_workbook(path, bundle):
    """Read-only QA of the saved artifact; export success alone is not proof."""
    from datetime import datetime
    import zipfile
    from io import BytesIO
    from openpyxl import load_workbook
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError('Workbook ZIP integrity failure')
    # Read via memory so failed QA cannot leave Windows ZIP-member file locks.
    workbook = load_workbook(BytesIO(Path(path).read_bytes()), read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ['Overview', 'Observations', 'Series coverage']:
            raise ValueError('Unexpected workbook sheets')
        for sheetname, rows in [('Observations', bundle['observations']), ('Series coverage', bundle['series_coverage'])]:
            sheet = workbook[sheetname]
            keys = list(rows[0])
            iterator = sheet.iter_rows(values_only=True)
            if list(next(iterator)) != keys:
                raise ValueError('Saved workbook column schema mismatch')
            count = 0
            for values, expected in zip(iterator, rows, strict=True):
                count += 1
                for key, saved in zip(keys, values, strict=True):
                    value = expected[key]
                    if key == 'ds':
                        equal = isinstance(saved, (date, datetime)) and saved.strftime('%Y-%m-%d') == value
                    elif key == 'y':
                        equal = isinstance(saved, (int, float)) and math.isclose(saved, float(value), rel_tol=1e-12, abs_tol=1e-8)
                    elif isinstance(value, (list, dict)):
                        equal = saved == json.dumps(value, ensure_ascii=False, separators=(',', ':'))
                    else:
                        equal = saved == value or (value == '' and saved is None)
                    if not equal:
                        raise ValueError(f'Saved workbook value mismatch: {sheetname}/{count+1}/{key}')
        if workbook['Overview']['B5'].value != len(bundle['observations']) or workbook['Overview']['B6'].value != len(bundle['closed_month_observations']):
            raise ValueError('Workbook overview count mismatch')
    finally:
        workbook.close()
    return True


def write_new_package(bundle, destination, root, node=None, exporter=None, workbook_source=None, source_renderer_returncode=None):
    """Atomic new-directory export. Existing destinations are never overwritten."""
    validate(bundle)
    destination, root = Path(destination).resolve(), Path(root).resolve()
    if destination.exists():
        raise FileExistsError('Use a new version; existing package is preserved')
    if not destination.parent.is_dir():
        raise ValueError('Destination parent must already exist')
    if workbook_source is not None and (node is not None or exporter is not None):
        raise ValueError('Use a fresh export or an independently verified saved artifact, not both')
    if workbook_source is not None:
        verify_workbook(workbook_source, bundle)
    tracked = {**bundle['source_hashes'], **bundle['publisher_hashes']}
    if any(sha(root / p) != h for p, h in tracked.items()):
        raise ValueError('Source snapshot changed since build')
    # Bundles cannot omit or change source observations while retaining audit claims.
    rebuilt = build(root, bundle['as_of'])
    for table in ['observations', 'closed_month_observations', 'chronological_partitions_proposed',
                  'observation_provenance', 'series_coverage']:
        if bundle[table] != rebuilt[table]:
            raise ValueError(f'Bundle changed relative to audited sources: {table}')
    with tempfile.TemporaryDirectory(prefix='.historical-staging-', dir=destination.parent) as stage:
        output = Path(stage) / 'output'; output.mkdir()
        tables = [('historical_observations', 'observations', OBSERVATION_FIELDS),
                  ('closed_month_observations', 'closed_month_observations', OBSERVATION_FIELDS),
                  ('chronological_partitions.proposed', 'chronological_partitions_proposed', SPLIT_FIELDS),
                  ('observation_provenance', 'observation_provenance', LINEAGE_FIELDS),
                  ('series_coverage', 'series_coverage', list(bundle['series_coverage'][0]))]
        for filename, table, fields in tables:
            _csv(output / f'{filename}.csv', bundle[table], fields)
        if (node is None) != (exporter is None):
            raise ValueError('Both Node and exporter are required for XLSX')
        exporter_returncode = None
        if workbook_source is not None:
            shutil.copyfile(workbook_source, output / 'historical_observations.xlsx')
            verify_workbook(output / 'historical_observations.xlsx', bundle)
            exporter_returncode = source_renderer_returncode
        if exporter:
            payload = Path(stage) / 'payload.json'
            payload.write_text(json.dumps(bundle, ensure_ascii=False, allow_nan=False), encoding='utf-8')
            process = subprocess.run([str(node), str(exporter), str(payload), str(output / 'historical_observations.xlsx'),
                                      str(root / 'work/historical_refinement_qa')], check=False)
            exporter_returncode = process.returncode
            if not (output / 'historical_observations.xlsx').is_file():
                raise ValueError(f'XLSX exporter produced no file (exit {process.returncode})')
            # The bundled native renderer can fail during shutdown after saving.
            # Recover only if the saved ZIP and every intended data cell reconcile.
            verify_workbook(output / 'historical_observations.xlsx', bundle)
        stats = {'version': VERSION, 'as_of': bundle['as_of'], 'observations': len(bundle['observations']),
                 'closed_month_rows': len(bundle['closed_month_observations']),
                 'open_or_future_month_rows': len(bundle['observations']) - len(bundle['closed_month_observations']),
                 'series_count': len(bundle['series_coverage']),
                 'proposed_split_series_count': len({r['series_id'] for r in bundle['chronological_partitions_proposed']}),
                 'proposed_split_rows': len(bundle['chronological_partitions_proposed']),
                 'provenance_rows': len(bundle['observation_provenance']),
                 'model_training_authorized': False, 'synthetic_rows_generated': 0,
                 'xlsx_exporter_returncode': exporter_returncode,
                 'xlsx_saved_cell_validation': 'PASS' if exporter or workbook_source else 'not requested',
                 'xlsx_renderer_warning': 'Native exporter returned nonzero; saved ZIP/schema/all data cells independently verified'
                                          if exporter_returncode not in (None, 0) else None,
                 'scope': 'Historical index observations and monthly aggregates of observed weekly prices; not total event costs',
                 'source_hashes': bundle['source_hashes'], 'publisher_hashes': bundle['publisher_hashes'],
                 'code_hashes': {str(p.relative_to(root)).replace('\\', '/'): sha(p) for p in
                                 [root / 'component4/historical_refinement.py', root / 'component4/historical_audit.py',
                                  root / 'scripts/research_pipeline/refine_historical_observations.py']},
                 'validation': 'PASS: schema, source arithmetic, origin, chronology and preservation checks',
                 'files': {p.name: {'sha256': sha(p), 'bytes': p.stat().st_size} for p in sorted(output.iterdir())}}
        if exporter:
            exporter_path = Path(exporter).resolve()
            stats['code_hashes'][str(exporter_path.relative_to(root)).replace('\\', '/')] = sha(exporter_path)
        (output / 'manifest.json').write_text(json.dumps(stats, indent=2, ensure_ascii=False, allow_nan=False) + '\n', encoding='utf-8')
        # Verify serialized tables, not just pre-export in-memory values.
        for filename, table, fields in tables:
            saved = read_csv(output / f'{filename}.csv')
            expected = [{k: '' if row[k] is None else json.dumps(row[k], ensure_ascii=False, separators=(',', ':'))
                         if isinstance(row[k], (list, dict)) else str(int(row[k]))
                         if isinstance(row[k], bool) else str(row[k]) for k in fields} for row in bundle[table]]
            if saved != expected:
                raise ValueError(f'Serialized table mismatch: {filename}')
        if any(sha(root / p) != h for p, h in tracked.items()):
            raise ValueError('Source snapshots changed during export')
        output.rename(destination)
    return stats

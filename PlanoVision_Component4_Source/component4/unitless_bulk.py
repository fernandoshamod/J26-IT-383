"""Separately authorized unitless expansion. No implicit pilot-to-bulk approval.

The frozen diagnostic calculator is reused unchanged. Root baselines are drawn
without replacement; overlapping interventions are grouped transitively.
Shared templates support conditional simulator evaluation only.
"""
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import date, timedelta
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import platform
import subprocess
import tempfile

import numpy as np
from .engine import FORMULA_IDS, fingerprint
from .unitless_pilot import validate as validate_pilot
from .unitless_experiment import _calculate_unitless
from .dataset_design import assign_groups, audit_split_leakage, split_components

ROOT = Path(__file__).resolve().parents[1]
VERSION = 'unitless-bulk.0.1'
INPUTS = ('guests', 'event_hours', 'stage_width_m', 'stage_length_m', 'budget_score', 'retained_receipts_score')
ORIGIN = 'SYNTHETIC_HYPOTHETICAL_UNITLESS_NOT_MARKET_DATA'
CODE_FILES = ('component4/unitless_bulk.py', 'component4/unitless_experiment.py',
              'component4/unitless_pilot.py', 'component4/engine.py', 'component4/dataset_design.py',
              'component4/diagnostic_design.py', 'work/minimal_pilot_authoring/export_unitless.mjs',
              'work/minimal_pilot_authoring/export_unitless_bulk.mjs',
              'scripts/research_pipeline/execute_unitless_bulk.py')


def hashes():
    return {name: hashlib.sha256((ROOT / name).read_bytes()).hexdigest() for name in CODE_FILES}


def validate(c):
    if not isinstance(c, dict) or c.get('version') != VERSION or c.get('cost_unit') != 'dimensionless_score' or c.get('currency') is not None:
        raise ValueError('explicit_version_and_unitless_contract_required')
    m = c['financial_model']
    validate_pilot(m)
    if fingerprint(m) != c.get('financial_model_sha256'):
        raise ValueError('financial_model_hash_mismatch')
    columns = c.get('master_columns')
    if not isinstance(columns, list) or len(columns) != 46 or len(set(columns)) != 46 or c.get('master_schema_sha256') != fingerprint(columns):
        raise ValueError('frozen_46_column_schema_required')
    s = c['sampling']
    for k in ('roots', 'states_per_root', 'replications_per_state', 'rows', 'seed', 'attempt_cap_per_root'):
        if type(s.get(k)) is not int or s[k] < (0 if k == 'seed' else 1):
            raise ValueError('positive_integer_required:' + k)
    if s['states_per_root'] != 3 or s['replications_per_state'] != 2 or s['rows'] != s['roots'] * 6:
        raise ValueError('three_states_two_replications_count_required')
    q = s.get('event_root_quotas')
    if not isinstance(q, dict) or set(q) != {'Wedding', 'Corporate'} or any(type(v) is not int or v < 1 for v in q.values()) or sum(q.values()) != s['roots']:
        raise ValueError('event_quotas_mismatch')
    if s.get('states') != ['baseline', 'guest_minus_one', 'guest_plus_one'] or s.get('root_policy') != 'unique_baselines_uniform_discrete_rejection_sampling':
        raise ValueError('explicit_sampling_design_required')
    start, end = date.fromisoformat(s['event_date_start']), date.fromisoformat(s['event_date_end'])
    if start > end:
        raise ValueError('date_bounds_invalid')
    if c.get('split_policy') != {'protocol': 'conditional_in_simulator', 'fractions': [0.6, 0.2, 0.2], 'tolerance': 0.05}:
        raise ValueError('explicit_conditional_split_required')
    if c.get('record_kind') not in ('hypothetical_unitless_bulk_proposal', 'mock_unitless_bulk_test'):
        raise ValueError('origin_marker_required')
    criteria = c.get('acceptance_criteria', {})
    if criteria != {'invalid_financial_rows': 0, 'invalid_spatial_rows': 0, 'missing_provenance_rows': 0,
                    'duplicate_ids': 0, 'distinct_baselines': s['roots'], 'unseen_template_evaluation': False,
                    'economic_realism_claim': False, 'class_balancing': False}:
        raise ValueError('acceptance_criteria_mismatch')
    if c.get('field_semantics_version') != 'unitless-master-mapping.0.1':
        raise ValueError('explicit_unitless_master_mapping_required')
    return True


def blockers(c, a=None):
    try:
        validate(c)
    except (ValueError, TypeError, KeyError) as exc:
        return ['invalid_configuration:' + str(exc)]
    a = a if isinstance(a, dict) else {}
    result = []
    if c['record_kind'] != 'hypothetical_unitless_bulk_proposal':
        result.append('mock_forbidden')
    expected = {'explicit_bulk_authorization': True, 'configuration_sha256': fingerprint(c),
                'code_sha256': hashes(), 'runtime': {'python': platform.python_version(), 'numpy': np.__version__},
                'scenario_count': c['sampling']['rows'], 'event_types': ['Wedding', 'Corporate']}
    for key, value in expected.items():
        if a.get(key) != value or (key == 'scenario_count' and type(a.get(key)) is not int):
            result.append('missing_or_mismatched:' + key)
    for key in ('authorizer', 'date', 'evidence_ref', 'approval_rationale'):
        if not isinstance(a.get(key), str) or not a[key].strip():
            result.append('missing:' + key)
    # Evidence is a real saved human decision, not a record manufactured by this module.
    evidence = ROOT / str(a.get('evidence_ref', '__missing__'))
    if not evidence.is_file() or a.get('evidence_sha256') != hashlib.sha256(evidence.read_bytes()).hexdigest():
        result.append('human_authorization_evidence_missing_or_changed')
    elif evidence.suffix == '.json':
        try:
            decision = json.loads(evidence.read_text(encoding='utf-8'))
            if (decision.get('configuration_sha256') != fingerprint(c) or decision.get('scenario_count') != c['sampling']['rows']
                    or decision.get('explicit_bulk_authorization') is not True or not decision.get('user_statement')):
                result.append('human_decision_scope_mismatch')
        except (ValueError, OSError):
            result.append('invalid_human_decision')
    else:
        result.append('machine_readable_human_decision_required')
    for key in ('pilot_review', 'test_receipt'):
        reference = c.get(key, {}) if key == 'pilot_review' else a.get(key, {})
        path = ROOT / str(reference.get('path', '__missing__'))
        if not path.is_file() or reference.get('sha256') != hashlib.sha256(path.read_bytes()).hexdigest():
            result.append('missing_or_changed:' + key)
        else:
            receipt = json.loads(path.read_text())
            if key == 'pilot_review' and (receipt.get('internal_status') != 'PASS' or receipt.get('configuration_sha256') != c['financial_model_sha256']):
                result.append('pilot_scope_or_status_mismatch')
            if key == 'test_receipt' and (receipt.get('exit_code') != 0 or receipt.get('code_sha256') != hashes() or receipt.get('configuration_sha256') != fingerprint(c)):
                result.append('test_receipt_mismatch')
    return sorted(set(result))


def _sample(c, mock=False):
    """Internal only; public persistent path always checks separate authorization."""
    validate(c)
    m, s = c['financial_model'], c['sampling']
    origin = 'MOCK_UNITLESS_TEST_NOT_MARKET_DATA' if mock else ORIGIN
    config_hash = fingerprint(c)
    planning, outcomes, lines, provenance, latent, rejected = [], [], [], [], [], []
    unique = set()
    events = [e for e in ('Wedding', 'Corporate') for _ in range(s['event_root_quotas'][e])]
    for index, event in enumerate(events):
        family = 'UBF-' + fingerprint({'config': config_hash, 'root': index})[:24]
        p = m['profiles'][event]
        rng = np.random.Generator(np.random.PCG64(np.random.SeedSequence(s['seed'], spawn_key=(index, 0))))
        for attempt in range(s['attempt_cap_per_root']):
            base = {k: p[k][int(rng.integers(len(p[k])))] for k in INPUTS if k != 'guests'}
            base['guests'] = int(rng.integers(p['guest_min'] + 1, p['guest_max']))
            signature = fingerprint({'event': event, 'inputs': base})
            if signature in unique:
                rejected.append({'family_id': family, 'attempt': attempt, 'reason': 'duplicate_baseline'})
                continue
            variants = [{**base, 'guests': base['guests'] + offset} for offset in (0, -1, 1)]
            try:
                results = [_calculate_unitless(m, event, x, origin) for x in variants]
            except ValueError as exc:
                rejected.append({'family_id': family, 'attempt': attempt, 'reason': str(exc)})
                continue
            unique.add(signature)
            break
        else:
            raise ValueError('root_attempt_cap_exhausted')
        start, end = date.fromisoformat(s['event_date_start']), date.fromisoformat(s['event_date_end'])
        event_date = (start + timedelta(days=int(rng.integers((end - start).days + 1)))).isoformat()
        for state, (inputs, r) in enumerate(zip(variants, results)):
            for rep in range(2):
                sid = 'UBS-' + fingerprint({'family': family, 'state': state, 'replication': rep})[:24]
                planning.append({'scenario_id': sid, 'family_id': family, 'event_type': event, 'event_date': event_date,
                                 'planning_state': s['states'][state], 'outcome_replication': rep, **inputs,
                                 'planned_cost_score': r['planned_cost_score'], 'reserve_score': r['reserve_score'],
                                 'planned_profit_score': r['planned_profit_score'], 'planned_margin_points': r['margin_points'],
                                 'used_area_m2': r['used_area_m2'], 'usable_area_m2': r['usable_area_m2'], 'staff': r['staff'],
                                 'tables': r['tables'], 'paid_hours_per_staff': r['paid_hours_per_staff'], 'seats': inputs['guests'],
                                 'planned_budget_breach': r['planned_budget_breach'],
                                 'operating_allowance_score': inputs['budget_score'] - r['reserve_score'],
                                 'planned_operating_allowance_breach': r['planned_cost_score'] > inputs['budget_score'] - r['reserve_score']})
                provenance.append({'scenario_id': sid, 'family_id': family,
                                   'planning_fingerprint': fingerprint({'event': event, 'inputs': inputs}),
                                   'venue_template_id': m['venue_template_id'], 'vendor_template_id': m['vendor_template_id'],
                                   'source_dependency_ids': m['source_dependency_ids'], 'configuration_sha256': config_hash,
                                   'financial_model_sha256': c['financial_model_sha256'], 'source_snapshot_sha256': m['source_snapshot_sha256'],
                                   'assumption_ids': list(FORMULA_IDS), 'origin': origin, 'geometry_origin': m['geometry_origin'],
                                   'cost_unit': 'dimensionless_score', 'currency': None, 'is_observed_event': False,
                                   'is_trained_prediction': False, 'event_date_origin': 'experimental_uniform_date_no_seasonal_price_effect'})
    # Split solely from planning provenance, BEFORE drawing any realized outcomes.
    policy = c['split_policy']
    assignments = assign_groups(provenance, policy['protocol'], s['seed'], tuple(policy['fractions']), policy['tolerance'])
    groups = split_components(provenance, policy['protocol'])
    for r in provenance:
        r.update(split=assignments[r['scenario_id']], split_group_id=groups[r['family_id']])
    for index, event in enumerate(events):
        for row in planning[index * 6:(index + 1) * 6]:
            inputs = {k: row[k] for k in INPUTS}
            r = _calculate_unitless(m, event, inputs, origin)
            rng = np.random.Generator(np.random.PCG64(np.random.SeedSequence(s['seed'], spawn_key=(index, 1, row['outcome_replication']))))
            o = m['outcome_proposal']
            cost_factor = float(rng.uniform(*o['ordinary_common_cost_multiplier']))
            receipt_factor = float(rng.uniform(*o['independent_receipt_multiplier']))
            injected = bool(rng.random() < o['injection_prevalence'])
            target = o['eligible_services'][int(rng.integers(len(o['eligible_services'])))] if injected else None
            multiplier = float(rng.uniform(*o['injection_multiplier'])) if injected else 1.0
            actual = {k: v * cost_factor * (multiplier if k == target else 1) for k, v in r['line_items_score'].items()}
            total = math.fsum(actual.values())
            receipts = inputs['retained_receipts_score'] * receipt_factor
            profit = receipts - total
            margin = 100 * profit / receipts if receipts else None
            sid = row['scenario_id']
            outcomes.append({'scenario_id': sid, 'simulated_cost_score': total, 'simulated_receipts_score': receipts,
                             'simulated_profit_score': profit, 'simulated_margin_points': margin,
                             'simulated_budget_overrun': total > inputs['budget_score'],
                             'simulated_profit_margin_risk': margin < m['minimum_margin_points'] if margin is not None else None,
                             'simulated_vendor_cost_anomaly': injected, 'label_origin': 'SIMULATOR_OUTCOME_NOT_OBSERVED'})
            latent.append({'scenario_id': sid, 'cost_multiplier': cost_factor, 'receipt_multiplier': receipt_factor,
                           'injected_service': target, 'injection_multiplier': multiplier})
            for service, amount in r['line_items_score'].items():
                lines.append({'scenario_id': sid, 'service': service, 'quantity': r['quantities'][service],
                              'quantity_unit': m['quantity_units'][service], 'reference_quantity': m['reference_quantities'][service],
                              'experimental_coefficient': m['coefficients'][service], 'planned_score': amount,
                              'simulated_realized_score': actual[service], 'cost_unit': 'dimensionless_score'})
    bundle = {'planning': planning, 'outcomes': outcomes, 'line_items': lines, 'provenance': provenance,
              'latent_DO_NOT_USE_AS_FEATURES': latent, 'rejections': rejected, 'configuration': c,
              'manifest': {'configuration_sha256': config_hash, 'code_sha256': hashes(), 'origin': origin,
                           'seed': s['seed'], 'python': platform.python_version(), 'numpy': np.__version__,
                           'realism_validated': False, 'split_audit': audit_split_leakage(provenance, assignments, policy['protocol'])}}
    validate_bulk_bundle(bundle)
    return bundle


def validate_bulk_bundle(bundle):
    c = bundle['configuration']
    validate(c)
    plans = bundle['planning']
    model = c['financial_model']
    mapped = {}
    for key in ('planning', 'outcomes', 'provenance', 'latent_DO_NOT_USE_AS_FEATURES'):
        rows = bundle[key]
        mapped[key] = {r['scenario_id']: r for r in rows}
        if len(mapped[key]) != len(rows):
            raise ValueError('duplicate_ids:' + key)
    if any(set(rows) != set(mapped['planning']) for rows in mapped.values()):
        raise ValueError('mismatched_ids')
    by_id = defaultdict(list)
    for line in bundle['line_items']:
        by_id[line['scenario_id']].append(line)
    if set(by_id) != set(mapped['planning']):
        raise ValueError('ledger_scope_mismatch')
    def close(a, b):
        return type(a) in (int, float) and math.isfinite(a) and math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-12)
    for p in plans:
        sid = p['scenario_id']
        expected = _calculate_unitless(model, p['event_type'], {k: p[k] for k in INPUTS})
        ledger, o, latent = by_id[sid], mapped['outcomes'][sid], mapped['latent_DO_NOT_USE_AS_FEATURES'][sid]
        if len(ledger) != 7 or {r['service'] for r in ledger} != set(expected['line_items_score']):
            raise ValueError('double_charging_or_missing_service')
        shock = model['outcome_proposal']
        if not (shock['ordinary_common_cost_multiplier'][0] <= latent['cost_multiplier'] <= shock['ordinary_common_cost_multiplier'][1]
                and shock['independent_receipt_multiplier'][0] <= latent['receipt_multiplier'] <= shock['independent_receipt_multiplier'][1]):
            raise ValueError('shock_outside_approved_bounds')
        if latent['injected_service'] is None:
            if latent['injection_multiplier'] != 1:
                raise ValueError('uninjected_multiplier_invalid')
        elif latent['injected_service'] not in shock['eligible_services'] or not shock['injection_multiplier'][0] <= latent['injection_multiplier'] <= shock['injection_multiplier'][1]:
            raise ValueError('injection_scope_invalid')
        for line in ledger:
            k = line['service']
            factor = latent['cost_multiplier'] * (latent['injection_multiplier'] if k == latent['injected_service'] else 1)
            if (line['quantity'] != expected['quantities'][k] or line['quantity_unit'] != model['quantity_units'][k]
                    or not close(line['planned_score'], expected['line_items_score'][k])
                    or not close(line['simulated_realized_score'], line['planned_score'] * factor)):
                raise ValueError('line_item_dependency_mismatch')
        if (not close(p['planned_cost_score'], math.fsum(r['planned_score'] for r in ledger))
                or not close(o['simulated_cost_score'], math.fsum(r['simulated_realized_score'] for r in ledger))
                or not close(o['simulated_receipts_score'], p['retained_receipts_score'] * latent['receipt_multiplier'])
                or not close(o['simulated_profit_score'], o['simulated_receipts_score'] - o['simulated_cost_score'])):
            raise ValueError('financial_reconciliation_failed')
        for key, expected_key in [('reserve_score', 'reserve_score'), ('planned_profit_score', 'planned_profit_score'),
                                  ('used_area_m2', 'used_area_m2'), ('usable_area_m2', 'usable_area_m2'),
                                  ('staff', 'staff'), ('tables', 'tables'), ('paid_hours_per_staff', 'paid_hours_per_staff')]:
            if not close(p[key], expected[expected_key]):
                raise ValueError('planning_dependency_mismatch:' + key)
        if (expected['margin_points'] is None and p['planned_margin_points'] is not None) or (expected['margin_points'] is not None and not close(p['planned_margin_points'], expected['margin_points'])):
            raise ValueError('planned_margin_invalid')
        if p['seats'] != p['guests'] or not close(p['operating_allowance_score'], p['budget_score'] - expected['reserve_score']):
            raise ValueError('seats_or_contingency_invalid')
        if p['planned_budget_breach'] != expected['planned_budget_breach'] or p['planned_operating_allowance_breach'] != (p['planned_cost_score'] > p['operating_allowance_score']):
            raise ValueError('planned_flag_invalid')
        margin = 100 * o['simulated_profit_score'] / o['simulated_receipts_score'] if o['simulated_receipts_score'] else None
        if (margin is None and o['simulated_margin_points'] is not None) or (margin is not None and not close(o['simulated_margin_points'], margin)):
            raise ValueError('margin_invalid')
        if (o['simulated_budget_overrun'] != (o['simulated_cost_score'] > p['budget_score'])
                or o['simulated_profit_margin_risk'] != (margin < model['minimum_margin_points'] if margin is not None else None)
                or o['simulated_vendor_cost_anomaly'] != (latent['injected_service'] is not None)
                or o['label_origin'] != 'SIMULATOR_OUTCOME_NOT_OBSERVED'):
            raise ValueError('label_invalid')
        d = date.fromisoformat(p['event_date'])
        if not date.fromisoformat(c['sampling']['event_date_start']) <= d <= date.fromisoformat(c['sampling']['event_date_end']):
            raise ValueError('event_date_outside_scope')
    if len({p['family_id'] for p in plans}) != c['sampling']['roots']:
        raise ValueError('root_count_mismatch')
    baselines = {(p['event_type'], tuple(p[k] for k in INPUTS)) for p in plans if p['planning_state'] == 'baseline'}
    if len(baselines) != c['sampling']['roots']:
        raise ValueError('duplicate_baselines')
    if Counter(p['event_type'] for p in plans) != Counter({k: v * 6 for k, v in c['sampling']['event_root_quotas'].items()}):
        raise ValueError('event_scope_count_mismatch')
    for row in bundle['provenance']:
        if (row['configuration_sha256'] != fingerprint(c) or row['financial_model_sha256'] != c['financial_model_sha256']
                or row['is_observed_event'] or row['is_trained_prediction'] or row['cost_unit'] != 'dimensionless_score'):
            raise ValueError('provenance_mismatch')
        if (row['venue_template_id'] != model['venue_template_id'] or row['vendor_template_id'] != model['vendor_template_id']
                or row['source_dependency_ids'] != model['source_dependency_ids'] or row['assumption_ids'] != list(FORMULA_IDS)
                or row['source_snapshot_sha256'] != model['source_snapshot_sha256']):
            raise ValueError('source_ancestry_mismatch')
        p = mapped['planning'][row['scenario_id']]
        if row['planning_fingerprint'] != fingerprint({'event': p['event_type'], 'inputs': {k: p[k] for k in INPUTS}}):
            raise ValueError('planning_fingerprint_mismatch')
    grouped = defaultdict(list)
    for p in plans:
        grouped[p['family_id']].append(p)
    for family, rows in grouped.items():
        if len(rows) != 6 or {(p['planning_state'], p['outcome_replication']) for p in rows} != {(state, rep) for state in c['sampling']['states'] for rep in (0, 1)}:
            raise ValueError('family_intervention_scope_mismatch')
        baseline = next(p for p in rows if p['planning_state'] == 'baseline')
        for p in rows:
            delta = {'baseline': 0, 'guest_minus_one': -1, 'guest_plus_one': 1}[p['planning_state']]
            if (p['guests'] != baseline['guests'] + delta or any(p[k] != baseline[k] for k in INPUTS if k != 'guests')
                    or p['event_type'] != baseline['event_type'] or p['event_date'] != baseline['event_date']):
                raise ValueError('one_factor_intervention_mismatch')
            state = c['sampling']['states'].index(p['planning_state'])
            if p['scenario_id'] != 'UBS-' + fingerprint({'family': family, 'state': state, 'replication': p['outcome_replication']})[:24]:
                raise ValueError('stable_scenario_id_mismatch')
            companion = next(q for q in rows if q['planning_state'] == 'baseline' and q['outcome_replication'] == p['outcome_replication'])
            one = {k: v for k, v in mapped['latent_DO_NOT_USE_AS_FEATURES'][p['scenario_id']].items() if k != 'scenario_id'}
            two = {k: v for k, v in mapped['latent_DO_NOT_USE_AS_FEATURES'][companion['scenario_id']].items() if k != 'scenario_id'}
            if one != two:
                raise ValueError('paired_outcome_shock_mismatch')
    if len(plans) != c['sampling']['rows'] or len(bundle['line_items']) != len(plans) * 7:
        raise ValueError('bulk_counts_invalid')
    audit = audit_split_leakage(bundle['provenance'], {r['scenario_id']: r['split'] for r in bundle['provenance']}, c['split_policy']['protocol'])
    if audit['status'] != 'PASS':
        raise ValueError('split_leakage')
    return True


def generate_test(c):
    if c.get('record_kind') != 'mock_unitless_bulk_test':
        raise ValueError('test_only_marker_required')
    return _sample(c, mock=True)


def generate(c, authorization):
    reasons = blockers(c, authorization)
    if reasons:
        raise ValueError('bulk_execution_blocked:' + '|'.join(reasons))
    return _sample(c)


def master_rows(bundle, columns):
    """Preserve schema names; all monetary columns are explicitly dimensionless.

Quoted_* fields are experimental planned charges, not quotations. Actual_* fields
are simulated realizations, not observations. Required sidecars preserve meaning.
"""
    m = bundle['configuration']['financial_model']
    if columns != bundle['configuration']['master_columns']:
        raise ValueError('schema_order_not_frozen')
    outcomes = {o['scenario_id']: o for o in bundle['outcomes']}
    ledgers = defaultdict(dict)
    for r in bundle['line_items']:
        ledgers[r['scenario_id']][r['service']] = r['planned_score']
    rows = []
    for p in bundle['planning']:
        o, ledger = outcomes[p['scenario_id']], ledgers[p['scenario_id']]
        values = {'Scenario_ID': p['scenario_id'], 'Base_Event_ID': p['family_id'], 'Event_ID': p['family_id'],
                  'Is_Augmented': True, 'Data_Source': bundle['manifest']['origin'], 'Event_Type': p['event_type'],
                  'Event_Date': p['event_date'], 'Month': date.fromisoformat(p['event_date']).month,
                  'Event_Duration_Hours': p['event_hours'], 'Guest_Count': p['guests'],
                  'Venue_Area_m2': m['venue_width_m'] * m['venue_length_m'], 'Usable_Area_m2': p['usable_area_m2'],
                  'Seating_Count': p['seats'], 'Stage_Area_m2': p['stage_width_m'] * p['stage_length_m'],
                  'Space_Utilization': p['used_area_m2'] / p['usable_area_m2'],
                  'Quoted_Catering_Cost': ledger['catering'], 'Quoted_AV_Cost': ledger['av'],
                  'Quoted_Venue_Rental_Cost': ledger['venue_fixed'],
                  'Quoted_Equipment_Rental_Cost': ledger['chairs'] + ledger['tables'] + ledger['stage'],
                  'Staff_Count': p['staff'], 'Labor_Hours': p['paid_hours_per_staff'] * p['staff'],
                  'Quoted_Labor_Cost': ledger['labour'], 'Initial_Budget': p['budget_score'],
                  'Expected_Revenue': p['retained_receipts_score'], 'Contingency_Reserve': p['reserve_score'],
                  'Required_Min_Profit_Margin': m['minimum_margin_points'],
                  'Forecasted_Total_Cost': p['planned_cost_score'], 'Actual_Total_Cost': o['simulated_cost_score'],
                  'Actual_Revenue': o['simulated_receipts_score'], 'Budget_Utilization': o['simulated_cost_score'] / p['budget_score'],
                  'Profit': o['simulated_profit_score'], 'Profit_Margin': o['simulated_margin_points'],
                  'Cost_Per_Guest': o['simulated_cost_score'] / p['guests'],
                  'Cost_Variance': o['simulated_cost_score'] - p['planned_cost_score'],
                  'Budget_Overrun': o['simulated_budget_overrun'], 'Profit_Margin_Risk': o['simulated_profit_margin_risk'],
                  'Vendor_Cost_Anomaly': o['simulated_vendor_cost_anomaly']}
        if not set(values) <= set(columns) or len(columns) != 46 or len(set(columns)) != 46:
            raise ValueError('46_column_schema_mapping_mismatch')
        rows.append({k: values.get(k) for k in columns})
    return rows


def write(c, authorization, destination, node, columns, exporter=None):
    # Gate runs before any rows or output directories are generated.
    bundle = generate(c, authorization)
    destination = Path(destination).resolve()
    if destination.exists() or not destination.parent.is_dir():
        raise ValueError('new_destination_in_existing_parent_required')
    bundle['master'] = master_rows(bundle, columns)
    validate_bulk_bundle(bundle)
    with tempfile.TemporaryDirectory(prefix='.unitless-bulk-staging-', dir=destination.parent) as temp:
        stage = Path(temp) / 'output'
        stage.mkdir()
        tables = {'master_event_financial_dataset': bundle['master'], 'master_scenario_provenance': bundle['provenance'],
                  'master_financial_line_items': bundle['line_items'], 'planning': bundle['planning'],
                  'outcomes': bundle['outcomes'], 'latent_DO_NOT_USE_AS_FEATURES': bundle['latent_DO_NOT_USE_AS_FEATURES'],
                  'rejections': bundle['rejections']}
        for name, rows in tables.items():
            with (stage / (name + '.csv')).open('w', newline='', encoding='utf-8') as handle:
                keys = list(rows[0]) if rows else ['family_id', 'attempt', 'reason']
                writer = csv.DictWriter(handle, fieldnames=keys)
                writer.writeheader()
                writer.writerows({k: json.dumps(v, sort_keys=True) if isinstance(v, (list, dict)) else v for k, v in r.items()} for r in rows)
        payload = stage / 'export_payload.json'
        payload.write_text(json.dumps(bundle, allow_nan=False), encoding='utf-8')
        if exporter:
            exporter(payload, stage / 'master_event_financial_dataset.xlsx')
        else:
            subprocess.run([str(node), str(ROOT / 'work/minimal_pilot_authoring/export_unitless_bulk.mjs'), str(payload),
                            str(stage / 'master_event_financial_dataset.xlsx')], check=True)
        if not (stage / 'master_event_financial_dataset.xlsx').is_file():
            raise ValueError('xlsx_export_missing')
        from openpyxl import load_workbook
        wb = load_workbook(stage / 'master_event_financial_dataset.xlsx', read_only=True, data_only=True)
        try:
            saved = list(wb['Master Data'].iter_rows(values_only=True))
            if list(saved[0]) != columns or len(saved) != len(bundle['master']) + 1:
                raise ValueError('saved_workbook_schema_or_count_mismatch')
            for saved_row, expected in zip(saved[1:], bundle['master']):
                for value, key in zip(saved_row, columns):
                    target = expected[key]
                    if type(value) in (int, float) and type(target) in (int, float):
                        equal = math.isclose(value, target, rel_tol=1e-12, abs_tol=1e-12)
                    elif hasattr(value, 'isoformat') and key == 'Event_Date':
                        equal = value.isoformat()[:10] == target
                    else:
                        equal = value == target
                    if not equal:
                        raise ValueError('saved_workbook_cell_mismatch:' + key)
        finally:
            wb.close()
        (stage / 'configuration.json').write_text(json.dumps(c, indent=2, allow_nan=False))
        (stage / 'authorization.json').write_text(json.dumps(authorization, indent=2))
        manifest = {**bundle['manifest'], 'rows': len(bundle['master']), 'families': c['sampling']['roots'],
                    'status': 'GENERATED_PENDING_FINAL_AUDIT', 'unseen_template_evaluation_supported': False,
                    'files': {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in stage.iterdir() if p.is_file()}}
        (stage / 'manifest.json').write_text(json.dumps(manifest, indent=2))
        if blockers(c, authorization):
            raise ValueError('authorization_changed_before_commit')
        os.rename(stage, destination)
    return destination

"""Isolated MOCK export QA, never an authorized research dataset."""
from copy import deepcopy
import hashlib
import json
import math
from pathlib import Path
import subprocess
import sys
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from component4.unitless_bulk import generate_test, master_rows, hashes


def main():
    dest = ROOT / 'work/unitless_bulk_export_verification'
    prior = ROOT / 'reports/unitless_bulk_authorization_v0.1/exporter_verification.json'
    if dest.exists():
        old = json.loads(prior.read_text()) if prior.is_file() else {}
        if not old or (old.get('status') != 'FAIL' and old.get('code_sha256', {}).get('work/minimal_pilot_authoring/export_unitless_bulk.mjs') == hashes()['work/minimal_pilot_authoring/export_unitless_bulk.mjs']):
            raise ValueError('refuse_overwrite_successful_fixture_QA')
        backup = dest / ('previous_verified_export.json' if old.get('status') == 'PASS' else 'initial_failed_verification.json')
        if not backup.exists():
            backup.write_bytes(prior.read_bytes())
    else:
        dest.mkdir()
    c = deepcopy(json.loads((ROOT / 'config/unitless_bulk_candidate.v0.1.proposed.json').read_text()))
    c['record_kind'] = 'mock_unitless_bulk_test'
    c['sampling'].update(roots=10, rows=60, event_root_quotas={'Wedding': 5, 'Corporate': 5})
    c['acceptance_criteria']['distinct_baselines'] = 10
    b = generate_test(c)
    b['master'] = master_rows(b, c['master_columns'])
    payload = dest / 'fixture_payload.json'
    payload.write_text(json.dumps(b, allow_nan=False), encoding='utf-8')
    output = dest / 'MOCK_TEST_ONLY_NOT_RESEARCH_DATA.xlsx'
    previous_xlsx = dest / 'previous_export.xlsx'
    if output.exists() and not previous_xlsx.exists():
        previous_xlsx.write_bytes(output.read_bytes())
    node = 'C:/Users/User/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe'
    process = subprocess.run([node, str(ROOT / 'work/minimal_pilot_authoring/export_unitless_bulk.mjs'), str(payload), str(output)],
                             capture_output=True, text=True, cwd=ROOT)
    previews = subprocess.run([node, str(ROOT / 'work/minimal_pilot_authoring/export_unitless_bulk.mjs'), str(payload), str(output), '--previews-only'],
                              capture_output=True, text=True, cwd=ROOT)
    checks = {}
    if output.is_file():
        wb = load_workbook(output, read_only=True, data_only=True)
        try:
            for sheet, key in [('Master Data', 'master'), ('Planning', 'planning'), ('Outcomes', 'outcomes'),
                               ('Line Items', 'line_items'), ('Provenance', 'provenance'),
                               ('Latent NOT FEATURES', 'latent_DO_NOT_USE_AS_FEATURES'), ('Rejections', 'rejections')]:
                rows = b[key]
                saved = list(wb[sheet].iter_rows(values_only=True))
                keys = list(rows[0]) if rows else ['family_id', 'attempt', 'reason']
                expected = [keys] + [[json.dumps(row[k], separators=(',', ':')) if isinstance(row[k], (dict, list)) else row[k] for k in keys] for row in rows]
                def equal(a, target):
                    if type(a) in (int, float) and type(target) in (int, float):
                        return math.isclose(a, target, rel_tol=1e-12, abs_tol=1e-12)
                    if hasattr(a, 'isoformat') and isinstance(target, str) and len(target) == 10:
                        return a.isoformat()[:10] == target
                    return a == target
                checks[sheet] = len(saved) == len(expected) and all(len(a) == len(target) and all(equal(x, y) for x, y in zip(a, target)) for a, target in zip(saved, expected))
        finally:
            wb.close()
    names = ['Overview', 'Master_Data', 'Planning', 'Outcomes', 'Line_Items', 'Provenance', 'Latent_NOT_FEATURES', 'Rejections']
    receipt = {'status': 'PASS' if process.returncode == 0 and len(checks) == 7 and all(checks.values()) and all((dest / (n + '.png')).is_file() for n in names) else 'FAIL',
               'scope': '60-row MOCK TEST ONLY. Zero research rows generated.', 'native_exit_code': process.returncode,
               'preview_native_exit_code': previews.returncode,
               'preview_warning': 'Preview process status is separate from saved export and numerical reconciliation.',
               'all_saved_cells': checks, 'code_sha256': hashes(), 'preview_directory': str(dest),
               'xlsx_sha256': hashlib.sha256(output.read_bytes()).hexdigest() if output.is_file() else None,
               'native_stdout': process.stdout[-2000:], 'native_stderr': process.stderr[-2000:]}
    (ROOT / 'reports/unitless_bulk_authorization_v0.1/exporter_verification.json').write_text(json.dumps(receipt, indent=2) + '\n')
    print(json.dumps(receipt, indent=2))
    if receipt['status'] != 'PASS':
        raise RuntimeError('Mock exporter QA failed; preserve fixture and diagnose before authorization')


if __name__ == '__main__':
    main()

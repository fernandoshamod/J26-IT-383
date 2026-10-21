"""Execute the exact user-approved 30-row unitless pilot; no bulk or model training."""
from collections import Counter
from copy import deepcopy
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import platform
import re
import subprocess
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
import numpy as np
from component4.engine import fingerprint, FORMULA_IDS
from component4.unitless_pilot import validate, code_hashes, execution_blockers, write, generate, validate_bundle

NODE = Path('C:/Users/User/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe')
DESTINATION = ROOT / 'data/pilot/unitless-v0.3-user-authorized-20261012'
REPORTS = ROOT / 'reports/unitless_pilot_execution_20261012'


def json_file(path):
    return json.loads(path.read_text(encoding='utf-8'))


def table(path):
    with path.open(newline='', encoding='utf-8') as f:
        return list(csv.DictReader(f))


def csv_expected(rows):
    return [{k: '' if v is None else json.dumps(v, sort_keys=True) if isinstance(v, (dict, list)) else str(v)
             for k, v in r.items()} for r in rows]


def main():
    if '--verify-existing' in sys.argv:
        verify_existing(json_file(DESTINATION / 'configuration.json'),
                        json_file(ROOT / 'config/unitless_pilot_authorization.v0.3.user.json'),
                        json_file(REPORTS / 'test_results.json'))
        return
    if DESTINATION.exists() or REPORTS.exists():
        raise ValueError('Existing pilot/receipt preserved; this command does not overwrite or generate another run.')
    c = json_file(ROOT / 'config/unitless_pilot_candidate.v0.3.proposed.json')
    evidence_path = ROOT / 'reports/unitless_pilot_user_authorization_evidence.json'
    evidence = json_file(evidence_path)
    revision = json_file(ROOT / 'reports/unitless_pilot_revision_manifest.json')
    validate(c)
    digest = fingerprint(c)
    if digest != evidence['configuration_sha256'] or digest != revision['configuration_sha256']:
        raise ValueError('Approved configuration hash mismatch')
    if code_hashes() != revision['code_sha256']:
        raise ValueError('Reviewed code snapshot changed')
    if evidence['scope'] != {'scenario_count': 30, 'event_types': ['Wedding', 'Corporate'],
                             'cost_unit': 'dimensionless_score', 'approval_kind': 'experiment_designer_acceptance',
                             'designer_decisions_approved': True, 'pilot_execution_authorized': True,
                             'supervisor_review_requirement_waived_by_user': True, 'qualified_review_performed': False,
                             'lkr_calibration_approved': False, 'bulk_generation_authorized': False,
                             'model_training_authorized': False}:
        raise ValueError('User approval scope changed')
    env = dict(os.environ, TEMP=str(ROOT / 'work'), TMP=str(ROOT / 'work'))
    test = subprocess.run([sys.executable, '-m', 'unittest', 'discover', '-s', 'tests', '-q'],
                          cwd=ROOT, env=env, capture_output=True, text=True)
    transcript = test.stdout + test.stderr
    match = re.search(r'Ran (\d+) tests? in ([0-9.]+)s', transcript)
    if test.returncode or not match:
        raise ValueError('Full test suite failed; generation blocked: ' + transcript)
    REPORTS.mkdir()
    (REPORTS / 'test_transcript.txt').write_text(transcript, encoding='utf-8')
    test_receipt = {'exit_code': test.returncode, 'tests_run': int(match[1]), 'elapsed_seconds': float(match[2]),
                    'transcript_sha256': hashlib.sha256(transcript.encode()).hexdigest(),
                    'code_sha256': code_hashes(), 'configuration_sha256': digest}
    (REPORTS / 'test_results.json').write_text(json.dumps(test_receipt, indent=2) + '\n')
    authorization = {
        'explicit_authorization': True, 'configuration_sha256': digest, 'code_sha256': code_hashes(),
        'runtime': {'python': platform.python_version(), 'numpy': np.__version__},
        'scenario_count': 30, 'event_types': ['Wedding', 'Corporate'],
        'authorizer': 'Human user in current chat; full name not supplied', 'date': evidence['recorded_date'],
        'evidence_ref': evidence_path.relative_to(ROOT).as_posix(),
        'evidence_sha256': hashlib.sha256(evidence_path.read_bytes()).hexdigest(),
        'approval_kind': 'experiment_designer_acceptance_with_qualified_review_waived',
        'qualified_review_performed': False, 'supervisor_review_requirement_waived_by_user': True,
        'institutional_approval_claimed': False, 'bulk_generation_authorized': False,
        'model_training_authorized': False,
        'approvals': {},
    }
    for name in (*FORMULA_IDS, 'numerical_parameters', 'methodology', 'sampling_split', 'acceptance_criteria', 'tests_passed'):
        authorization['approvals'][name] = {
            'status': 'approved', 'reviewer': 'Human user in current chat; full name not supplied',
            'role': 'experiment_designer_NOT_supervisor', 'date': evidence['recorded_date'],
            'rationale': 'User approved all choices for the reviewed 30-row unitless diagnostic experiment and explicitly waived supervisor review. No empirical validation or professional sign-off claimed.',
            'evidence_ref': authorization['evidence_ref'], 'configuration_sha256': digest,
            'approval_kind': 'hypothetical_experiment_only', 'qualified_review': False,
        }
    authorization['approvals']['tests_passed']['rationale'] = 'User accepts the diagnostic execution; the full software suite was rerun successfully before generation. See bound test receipt.'
    authorization['approvals']['tests_passed']['test_receipt'] = (REPORTS / 'test_results.json').relative_to(ROOT).as_posix()
    blockers = execution_blockers(c, authorization)
    if blockers:
        raise ValueError('Execution gate remains blocked:' + '|'.join(blockers))
    auth_path = ROOT / 'config/unitless_pilot_authorization.v0.3.user.json'
    if auth_path.exists():
        raise ValueError('Existing user authorization preserved')
    auth_path.write_text(json.dumps(authorization, indent=2) + '\n')
    DESTINATION.parent.mkdir(exist_ok=True)
    write(c, authorization, DESTINATION, NODE)
    verify_existing(c, authorization, test_receipt)


def verify_existing(c, authorization, test_receipt):
    digest = fingerprint(c)
    if execution_blockers(c, authorization):
        raise ValueError('Authorization/code/configuration changed before verification')
    replay = generate(c, authorization)
    validate_bundle(replay)
    manifest = json_file(DESTINATION / 'manifest.json')
    checks = {}
    for name in ('planning', 'outcomes', 'line_items', 'provenance', 'latent_DO_NOT_USE_AS_FEATURES', 'rejections'):
        actual = table(DESTINATION / (name + '.csv'))
        checks[name + '_exact_replay'] = actual == csv_expected(replay[name])
    checks['all_written_file_hashes'] = all(hashlib.sha256((DESTINATION / name).read_bytes()).hexdigest() == value for name, value in manifest['files'].items())
    checks['authorized_configuration_match'] = fingerprint(json_file(DESTINATION / 'configuration.json')) == digest
    checks['all_30_unique_ids'] = len({x['scenario_id'] for x in replay['planning']}) == 30
    checks['five_base_families'] = len({x['family_id'] for x in replay['planning']}) == 5
    checks['seven_disjoint_charges_per_row'] = len(replay['line_items']) == 210
    checks['group_split_audit'] = replay['manifest']['split_audit']['status'] == 'PASS'
    checks['truthful_origin'] = all(x['origin'] == 'SYNTHETIC_HYPOTHETICAL_UNITLESS_NOT_MARKET_DATA' and not x['is_observed_event'] and not x['is_trained_prediction'] for x in replay['provenance'])
    checks['realism_not_claimed'] = replay['manifest']['realism_validated'] is False
    # Read-only saved workbook QA; never authors via openpyxl.
    from openpyxl import load_workbook
    wb = load_workbook(DESTINATION / 'unitless_dataset.xlsx', read_only=True, data_only=True)
    try:
        for name, sheet in [('planning', 'planning'), ('outcomes', 'outcomes'), ('line_items', 'line_items'), ('provenance', 'provenance'), ('latent_DO_NOT_USE_AS_FEATURES', 'Latent NOT FEATURES'), ('rejections', 'rejections')]:
            rows = replay[name]
            saved = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
            keys = saved[0]
            required = set(rows[0]) if rows else {'family_id', 'attempt', 'reason'}
            checks[sheet + '_exact_header_set'] = len(keys) == len(set(keys)) and set(keys) == required
            if not checks[sheet + '_exact_header_set']:
                checks[sheet + '_all_saved_cells_reconcile'] = False
                continue
            # The JSON writer sorts keys. Compare by saved column names, not the
            # in-memory insertion order; do not rewrite any exported data.
            expected = [keys] + [[json.dumps(r[k], separators=(',', ':')) if isinstance(r[k], (dict, list)) else r[k] for k in keys] for r in rows]
            def same(a, b):
                if isinstance(a, bool) or isinstance(b, bool):
                    return type(a) is type(b) and a == b
                if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                    return math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-12)
                return a == b
            checks[sheet + '_all_saved_cells_reconcile'] = len(saved) == len(expected) and all(
                len(a) == len(b) and all(same(x, y) for x, y in zip(a, b)) for a, b in zip(saved, expected))
    finally:
        wb.close()
    distribution = {key: dict(Counter(str(x[key]) for x in replay['outcomes'])) for key in ('simulated_budget_overrun', 'simulated_profit_margin_risk', 'simulated_vendor_cost_anomaly')}
    result = {'status': 'PASS' if all(checks.values()) else 'FAIL', 'scope': 'Internal unitless pilot consistency ONLY; no real-world or qualified-review validation',
              'rows': 30, 'families': 5, 'planning_states_per_family': 3, 'outcome_replications_per_state': 2,
              'event_counts': dict(Counter(x['event_type'] for x in replay['planning'])),
              'line_items': 210, 'rejected_configuration_attempts': len(replay['rejections']),
              'class_distributions': distribution, 'checks': checks, 'configuration_sha256': digest,
              'qualified_review_performed': False, 'research_dataset_ready': False,
              'test_results': test_receipt, 'bulk_generation_authorized': False, 'model_training_authorized': False}
    prior = REPORTS / 'pilot_validation_report.json'
    retained = REPORTS / 'pilot_validation_report.initial_failed_check.json'
    if prior.exists() and json_file(prior)['status'] == 'FAIL' and not retained.exists():
        retained.write_bytes(prior.read_bytes())
    result['verification_note'] = 'Saved workbook columns are alphabetically ordered by the existing JSON serializer; all cells compared by header name with numeric tolerance 1e-12. Initial order-only checker failure is retained. No CSV rows, labels or financial rules changed. Any workbook presentation-only transform is recorded in the manifest.'
    prior.write_text(json.dumps(result, indent=2) + '\n')
    (REPORTS / 'pilot_execution_report.md').write_text(
        '# User-authorized unitless pilot execution\n\n'
        '30 rows generated: 18 Wedding, 12 Corporate; five families, three planning states and two outcome replications. '
        'Seven disjoint charges yield 210 line items. All numerical outputs are dimensionless, hypothetical and not LKR quotations.\n\n'
        'The human user approved the reviewed version 0.3 and explicitly waived supervisor review. '
        'Approval is experiment-designer acceptance, not supervisor/domain-expert or institutional approval. '
        'The earlier changes-requested record and blank authorization template remain unchanged as historical records.\n\n'
        'Canonical configuration SHA-256: `' + digest + '`. Code, runtime, seed, approvals and source snapshot are retained in the generated manifest.\n\n'
        'Internal pilot validation: ' + result['status'] + '. Check details and actual label counts are in pilot_validation_report.json. '
        'No class balancing, anomaly forcing, bulk generation or model training occurred. '
        'This is not a Stage 6 real-world realism PASS or a validated 7,500-row final dataset. '
        'The Stage 10 NOT READY archive is retained as a pre-pilot historical audit snapshot.\n', encoding='utf-8')
    if not all(checks.values()):
        raise ValueError('Pilot generated but validation FAILED; preserve outputs and investigate')
    print(json.dumps(result, indent=2))


if __name__ == '__main__':
    main()

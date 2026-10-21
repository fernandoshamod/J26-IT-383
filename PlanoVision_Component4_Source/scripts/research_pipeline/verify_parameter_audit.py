"""Independent read-only checks for the evidence audit artifacts."""
from pathlib import Path
import json, hashlib
from datetime import datetime
import openpyxl
from component4.engine import evaluate, PARAMS, RATE_UNITS

ROOT=Path(__file__).resolve().parents[2]
proposal=json.loads((ROOT/'config/simulator_parameters.proposed.json').read_text(encoding='utf-8'))
audit=json.loads((ROOT/'work/parameter_audit_tables.json').read_text(encoding='utf-8'))
for p,h in audit['snapshot'].items():
    assert hashlib.sha256((ROOT/p).read_bytes()).hexdigest()==h, 'Changed original: '+p
for kind in ['Wedding','Corporate']:
    profile=proposal['archetypes'][kind]
    assert all(profile['parameters'][k] is None for k in PARAMS)
    assert all(profile['rates'][k]['rate'] is None and profile['rates'][k]['tax_rate_points'] is None for k in RATE_UNITS)
    assert profile['parameter_approval']['status']=='pending'
    assert evaluate(proposal,kind,mode='live')['status']=='blocked'
assert proposal['pilot']['explicit_user_authorization'] is False
assert proposal['pilot']['scenario_count'] is None
assert len(proposal['public_price_candidates']['rates'])==20
expected_ids={'cbsl_prices_wages.csv':'S017','events.csv':'S015','inflation.csv':'S001;S002','retail_prices.csv':'S003','sl_procurement.csv':'S004','uk_service_awards.csv':'S008','usda_away.csv':'S010','usda_home_national.csv':'S018','usda_home_supplement.csv':'S018','venues.csv':'S013'}
assert {r[0]:r[1] for r in audit['Existing_Sources']['rows']}==expected_ids

price_file=ROOT/'reports/public_lkr_rate_evidence.xlsx'
check_file=ROOT/'reports/unresolved_parameter_checklist.xlsx'
prices=openpyxl.load_workbook(price_file,data_only=True)
checks=openpyxl.load_workbook(check_file,data_only=True)
for wb,counts in [(prices,{'Rates':20,'Source_Terms':9,'Unverified_Leads':6}),(checks,{'Parameters':108,'Controls':172,'Existing_Sources':10,'Master_Schema':46,'Assumption_Links':34})]:
    for name,count in counts.items():
        sheet=wb[name]
        assert len(sheet.tables)==1
        table=next(iter(sheet.tables.values()))
        assert int(table.ref.split(':')[-1].lstrip('ABCDEFGHIJKLMNOPQRSTUVWXYZ'))==count+1
        assert sheet.freeze_panes is not None
for i in range(2,22):
    assert isinstance(prices['Rates'].cell(i,13).value,datetime)
    assert prices['Rates'].cell(i,12).value is None
    assert prices['Rates'].cell(i,15).value=='Pending review'
    assert prices['Rates'].cell(i,19).value.startswith('https://')
assert prices['Rates']['E2'].value==6500
assert prices['Rates']['E12'].value is None
assert prices['Rates']['F12'].value==75 and prices['Rates']['G12'].value==80
assert prices['Source_Terms']['D17'].value==40250
assert all(checks['Parameters'].cell(i,8).value is None for i in range(2,110))
assert all(checks['Controls'].cell(i,4).value is None for i in range(2,174))
assert [checks['Review'].cell(i,2).value for i in range(6,13)]==[108,54,54,108,172,46,34]
formulas=openpyxl.load_workbook(price_file,data_only=False)
assert formulas['Source_Terms']['D17'].value=='=D15*(1+D16/100)'
for wb in [prices,checks]:
    for sheet in wb:
        for row in sheet:
            for cell in row:
                assert cell.data_type!='e',(sheet.title,cell.coordinate,cell.value)
print('Verified: 2 workbooks, 20 rates, 108 null parameter fields, 172 controls, 46 columns, 34 assumptions, preserved originals, blocked live profiles.')

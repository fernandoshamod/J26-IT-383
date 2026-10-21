"""Verify saved recovery artifacts and full software suite, without approving or generating datasets."""
import hashlib
import io
import json
from pathlib import Path
import unittest
from zipfile import ZipFile
import openpyxl

from component4.scenario_generator import pilot_execution_blockers
from component4.unitless_experiment import validate_unitless_config, unitless_pilot_blockers

ROOT = Path(__file__).resolve().parents[2]


def sha(p):
    return hashlib.sha256((ROOT / p).read_bytes()).hexdigest()


def read(p):
    return json.loads((ROOT / p).read_text(encoding='utf-8-sig'))


def main():
    expected = ['reports/master_dataset_blocker_register.xlsx', 'reports/formula_and_parameter_decisions.md',
                'reports/sampling_ancestry_review.md', 'reports/group_split_strategy.md',
                'reports/stage5_recovery_plan.md', 'reports/pilot_approval_decision_packet.md']
    assert all((ROOT / p).is_file() for p in expected)
    audit = read('work/dataset_recovery_audit.json')
    xlsx = expected[0]
    with ZipFile(ROOT / xlsx) as z:
        assert z.testzip() is None
        names = z.namelist()
        assert any(n.startswith('xl/tables/') for n in names)
    book = openpyxl.load_workbook(ROOT / xlsx, read_only=False, data_only=True)
    assert book.sheetnames == ['Review', 'Blockers']
    assert [book['Review'][f'B{i}'].value for i in range(6,12)] == [153,148,4,1,19,108]
    rows = list(book['Blockers'].iter_rows(min_row=2, max_row=154, values_only=True))
    assert len(rows) == 153 and len({r[0] for r in rows}) == 153
    assert [r[0] for r in rows] == [b['Blocker_ID'] for b in audit['blockers']]
    assert all(all(x is None for x in row[10:13]) for row in rows)
    assert not any(c.data_type == 'e' for s in book for row in s for c in row)
    assert book['Blockers'].freeze_panes is not None
    assert len(book['Blockers'].data_validations.dataValidation) == 1
    book.close()
    original = read('config/minimal_pilot_candidate.json')
    unitless = read('config/unitless_pilot.proposed.json')
    assert sha('config/minimal_pilot_candidate.json') == 'ce968556a717167c10a3b7bed406edd957f835a425d69984a3d8ed1080ad2e7b'
    assert len(pilot_execution_blockers(original)) == 167
    assert validate_unitless_config(unitless)
    assert unitless_pilot_blockers(unitless)
    assert not unitless['execution_enabled'] and not unitless['pilot_authorization']['explicit_authorization']
    assert unitless['frozen_configuration_sha256'] is None
    historical = read('reports/final_dataset_manifest.candidate.json')
    for p, h in historical['code_current_sha256'].items():
        assert sha(p) == h, f'Original code changed: {p}'
    for p, d in historical['preexisting_final_artifacts_preserved'].items():
        assert sha(p) == d['sha256'], f'Existing data changed: {p}'
    assert not (ROOT / 'data/pilot').exists()
    stream = io.StringIO()
    result = unittest.TextTestRunner(stream=stream, verbosity=1).run(unittest.defaultTestLoader.discover(str(ROOT / 'tests')))
    assert result.wasSuccessful() and result.testsRun >= 115 and not result.skipped
    verification = {'saved_workbook_zip_and_structure': 'PASS', 'summary_formula_cached_results': 'PASS',
                    'reviewer_fields_blank': True, 'visual_review': 'Both sheets inspected; register across identity/action/reviewer panels',
                    'summary_status_edit_and_restoration': 'PASS in authoring engine',
                    'authoring_exit_disposition': 'Authoring returned exit 1 after reporting successful export; independently reopened saved file passed checks',
                    'native_desktop_excel_tested': False, 'tests_run': result.testsRun,
                    'failures': len(result.failures), 'errors': len(result.errors), 'skipped': len(result.skipped),
                    'test_output': stream.getvalue(), 'original_code_and_data_unchanged': True,
                    'original_lkr_gate_codes': 167, 'unitless_proposal_structure': 'PASS',
                    'unitless_execution_gate': 'BLOCKED', 'unitless_blockers': unitless_pilot_blockers(unitless),
                    'pilot_rows_generated': 0, 'bulk_rows_generated': 0, 'stage6_pass_claim': False,
                    'deliverable_sha256': {p: sha(p) for p in expected}}
    (ROOT / 'work/dataset_recovery_verification.json').write_text(json.dumps(verification, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: v for k,v in verification.items() if k not in ('test_output','deliverable_sha256')}, indent=2))


if __name__ == '__main__':
    main()

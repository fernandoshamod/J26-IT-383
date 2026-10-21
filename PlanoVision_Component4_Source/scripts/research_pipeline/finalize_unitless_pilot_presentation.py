"""Promote verified presentation formatting without changing scenario data."""
import hashlib
import json
import os
from pathlib import Path
import shutil
import zipfile
from openpyxl import load_workbook  # Saved-cell reading only; no workbook authoring.

ROOT = Path(__file__).resolve().parents[2]
PILOT = ROOT / 'data/pilot/unitless-v0.3-user-authorized-20261012'
REPORTS = ROOT / 'reports/unitless_pilot_execution_20261012'
FORMATTED = ROOT / 'work/pilot_visual_qa_20261012_formatted/unitless_dataset.formatted.xlsx'


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    target = PILOT / 'unitless_dataset.xlsx'
    manifest_path = PILOT / 'manifest.json'
    manifest = json.loads(manifest_path.read_text())
    if 'presentation_transform' in manifest:
        raise ValueError('Presentation already finalized; no overwrite')
    a, b = load_workbook(target, read_only=True), load_workbook(FORMATTED, read_only=True)
    try:
        assert a.sheetnames == b.sheetnames
        for name in a.sheetnames:
            before, after = list(a[name].values), list(b[name].values)
            assert before == (after[:6] if name == 'Read first' else after), name
    finally:
        a.close(); b.close()
    backup = ROOT / 'work/pilot_visual_qa_20261012/unitless_dataset.original_export.xlsx'
    if backup.exists():
        raise ValueError('Original backup already exists; no overwrite')
    original_sha = sha(target)
    shutil.copy2(target, backup)
    shutil.copy2(FORMATTED, PILOT / 'unitless_dataset.presentation-staged.xlsx')
    os.replace(PILOT / 'unitless_dataset.presentation-staged.xlsx', target)
    manifest['files']['unitless_dataset.xlsx'] = sha(target)
    manifest['presentation_transform'] = {
        'scope': 'Column widths/wrapping/row heights/gridlines plus two explicit approval-scope notes; all six scenario tables cell-identical',
        'original_xlsx_sha256': original_sha, 'formatted_xlsx_sha256': sha(target),
        'formatter_sha256': sha(ROOT / 'work/minimal_pilot_authoring/inspect_authorized_pilot.mjs'),
        'original_backup': backup.relative_to(ROOT).as_posix(),
        'seven_sheet_previews_visually_reviewed': True,
        'renderer_process_exit_code': 1,
        'renderer_warning': 'Renderer returned nonzero after writing all seven previews and the saved workbook. All previews were inspected; all scenario cells independently match original export.',
    }
    manifest_path.write_text(json.dumps(manifest, indent=2) + '\n')
    (REPORTS / 'presentation_verification.json').write_text(json.dumps(manifest['presentation_transform'], indent=2) + '\n')
    print(json.dumps({'presentation_verified': True, 'all_scenario_cells_unchanged': True, 'formatted_sha256': sha(target)}))


if __name__ == '__main__':
    main()

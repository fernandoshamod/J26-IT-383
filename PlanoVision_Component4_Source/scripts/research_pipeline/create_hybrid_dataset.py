"""Prepare derivatives and package existing observations/scenarios; never generate events."""
import argparse
from collections import Counter, defaultdict
from datetime import date, datetime
import json
import math
import os
from pathlib import Path
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
import numpy as np
from openpyxl import load_workbook
from component4.engine import fingerprint
from component4.hybrid_dataset import (BASKET, TARGETS, FEATURES, FORBIDDEN, NODE_TYPES,
    NODE_FEATURES, EDGE_RULES, sha, read_csv, write_csv, historical_selection,
    chronological_audit, make_xgboost, graph_nodes, preserved_split_audit)

SYNTHETIC = ROOT / 'data/final/unitless-bulk-v0.1'
HISTORICAL = ROOT / 'data/refined/historical.v0.1-20261012'
EXPECTED_HASH = '2a94673fa623cae6aff278f44d103061129718697dbded14ef1e8fe42404b9bb'


def worksheet_records(sheet, identity_prefix=None):
    iterator = sheet.iter_rows(values_only=True)
    header = list(next(iterator))
    records = [{k: value.date().isoformat() if isinstance(value, datetime) else value
             for k, value in zip(header, row) if k} for row in iterator if any(v is not None for v in row)]
    if identity_prefix:
        identity = 'Evidence_ID' if identity_prefix == 'E' else 'Source_ID'
        records = [r for r in records if re.fullmatch(identity_prefix + r'\d{2,3}', str(r.get(identity, '')))]
    return records


def prepare(stage, test_receipt):
    protected = {str(p.relative_to(ROOT)).replace('\\', '/'): sha(p) for p in SYNTHETIC.iterdir() if p.is_file()}
    manifest = json.loads((SYNTHETIC / 'manifest.json').read_text())
    if any(sha(SYNTHETIC / p) != h for p, h in manifest['files'].items()):
        raise ValueError('original_synthetic_checksum_changed')
    bundle = json.loads((SYNTHETIC / 'export_payload.json').read_text())
    if fingerprint(bundle['configuration']) != EXPECTED_HASH:
        raise ValueError('frozen_configuration_mismatch')
    prior = json.loads((ROOT / 'reports/unitless_bulk_execution_20261012/validation_results.json').read_text())
    if prior['status'] != 'PASS_INTERNAL_UNITLESS_DATASET_AUDIT' or prior['configuration_sha256'] != EXPECTED_HASH:
        raise ValueError('missing_existing_internal_validation')
    hmanifest = json.loads((HISTORICAL / 'manifest.json').read_text())
    for name, record in hmanifest['files'].items():
        if sha(HISTORICAL / name) != record['sha256']:
            raise ValueError('changed_historical_export:' + name)
    for field in ('source_hashes', 'publisher_hashes'):
        if any(sha(ROOT / p) != value for p, value in hmanifest[field].items()):
            raise ValueError('changed_historical_source:' + field)
    selected, links = historical_selection(read_csv(HISTORICAL / 'closed_month_observations.csv'), read_csv(HISTORICAL / 'observation_provenance.csv'))
    ids = {r['observation_id'] for r in selected}
    raw_ids = {r['source_record_id'] for r in links}
    prices = [r for r in read_csv(ROOT / 'data/processed/retail_prices.csv') if r['source_record_id'] in raw_ids]
    indices = [r for r in read_csv(ROOT / 'data/processed/inflation.csv') if r['source_record_id'] in raw_ids]
    if len(prices) + len(indices) != len(raw_ids):
        raise ValueError('raw_observation_duplicate_or_missing')
    monthly = {r['observation_id']: r for r in selected}
    by_source = {r['source_record_id']: r for r in links}
    for row in prices:
        link = by_source[row['source_record_id']]
        if float(row['current_price_lkr']) != float(link['source_observed_value']) or row['parse_status'] != 'Parsed':
            raise ValueError('raw_price_lineage_mismatch')
        row.update(currency='LKR', unit='LKR/' + row['unit_raw'], series_id=monthly[link['observation_id']]['series_id'],
                   date_precision='publisher_week_label_not_exact_day', synthetic_origin=0, observed_event_financial_outcome=0,
                   publisher_url=link['publisher_url'], publisher_sha256=link['publisher_sha256'])
    for row in indices:
        link = by_source[row['source_record_id']]
        if float(row['index_value']) != float(link['source_observed_value']):
            raise ValueError('index_lineage_mismatch')
        row.update(currency=None, unit=link['source_unit'], synthetic_origin=0,
                   date_precision='monthly_period_key_not_publication_date', observed_event_financial_outcome=0,
                   publisher_url=link['publisher_url'], publisher_sha256=link['publisher_sha256'])
    prices.sort(key=lambda r: (r['series_id'], r['reported_week']))
    indices.sort(key=lambda r: (r['source_id'], r['month']))
    for rows, columns in ((prices, ('source_id', 'item_raw', 'unit_raw', 'reported_week')), (indices, ('source_id', 'index_type', 'base_year', 'month'))):
        if len({tuple(r[k] for k in columns) for r in rows}) != len(rows):
            raise ValueError('duplicate_real_measurement_key')
    write_csv(stage / 'prepared/real/weekly_retail_observations.csv', prices)
    write_csv(stage / 'prepared/real/monthly_inflation_observations.csv', indices)
    write_csv(stage / 'prepared/real/observation_provenance.csv', links)
    prophet = [{**r, 'currency': 'LKR' if r['source_id'] == 'S003' else None,
                'period_precision': 'month', 'event_total_cost_observation': 0} for r in selected]
    partitions = [r for r in read_csv(HISTORICAL / 'chronological_partitions.proposed.csv') if r['observation_id'] in ids]
    coverage = chronological_audit(prophet, partitions)
    write_csv(stage / 'prepared/prophet/prophet_cost_timeseries.csv', prophet)
    write_csv(stage / 'prepared/prophet/chronological_split_manifest.csv', partitions)
    write_csv(stage / 'prepared/prophet/series_coverage.csv', coverage)
    splitmap = {r['observation_id']: r['partition'] for r in partitions}
    for split in ('train', 'validation', 'test'):
        write_csv(stage / f'prepared/prophet/{split}.csv', [r for r in prophet if splitmap[r['observation_id']] == split])
    workbook = load_workbook(ROOT / 'reports/public_lkr_rate_evidence.xlsx', read_only=True, data_only=True)
    try:
        rates = worksheet_records(workbook['Rates'], 'E')
        terms = worksheet_records(workbook['Source_Terms'], 'P')
    finally:
        workbook.close()
    for row in rates:
        row.update(Currency='LKR', Record_Kind='published_advertised_rate_not_observed_transaction',
                   Approved_For_Simulator=False, Original_Evidence_File_SHA256=sha(ROOT / 'reports/public_lkr_rate_evidence.xlsx'))
    write_csv(stage / 'prepared/benchmarks/published_lkr_rate_evidence.csv', rates)
    write_csv(stage / 'prepared/benchmarks/provider_terms.csv', terms)
    downloads = json.loads((ROOT / 'outputs/01a12103-7a09-7331-bb93-9e24a6f05005/download_manifest.json').read_text())
    procurement_source = next(r for r in downloads if r['source_id'] == 'S004' and r['status'] == 'Downloaded')
    pdf = ROOT / 'data/raw/publishers/S004' / procurement_source['name']
    if sha(pdf) != procurement_source['sha256']:
        raise ValueError('procurement_publisher_changed')
    procurement = read_csv(ROOT / 'data/processed/sl_procurement.csv')
    for row in procurement:
        row.update(currency='LKR', unit='LKR million as published', publisher_url=procurement_source['url'],
                   publisher_sha256=procurement_source['sha256'], record_kind='overlapping_aggregate_procurement_context',
                   event_quote=False, model_input=False, warning='Do not sum overlapping periods or divide to invent per-event prices')
    write_csv(stage / 'prepared/benchmarks/sl_procurement_context.csv', procurement)
    provenance = bundle['provenance']
    split_audit = preserved_split_audit(provenance)
    xgb = make_xgboost(bundle['master'], provenance)
    write_csv(stage / 'prepared/xgboost/xgboost_risk_dataset.csv', xgb)
    splitrows = [{'Scenario_ID': r['scenario_id'], 'Base_Event_ID': r['family_id'], 'planning_fingerprint': r['planning_fingerprint'],
                  'split_group_id': r['split_group_id'], 'split': r['split'], 'configuration_sha256': r['configuration_sha256']} for r in provenance]
    write_csv(stage / 'prepared/xgboost/grouped_split_manifest.csv', splitrows)
    for split in ('train', 'validation', 'test'):
        write_csv(stage / f'prepared/xgboost/{split}.csv', [r for r in xgb if r['split'] == split])
    spec = {'version': 'hybrid-features.1.0', 'predictors': list(FEATURES), 'metadata_not_predictors': list(xgb[0])[:5],
            'targets': list(TARGETS), 'forbidden_inputs': sorted(FORBIDDEN), 'transformations_fitted': False,
            'encoding': 'fixed Wedding/Corporate one-hot, not data-fitted', 'planned_forecast_origin': 'frozen deterministic rule engine, not ML forecast',
            'unit': 'dimensionless_score for financial fields', 'labels': 'synthetic_outcomes_not_observed_events',
            'vendor_anomaly_limit': 'Injection is random and independent of planning features; meaningful prediction is not expected from this feature set.'}
    (stage / 'prepared/xgboost/feature_specification.json').write_text(json.dumps(spec, indent=2))
    dictionary = [{'dataset': 'XGBoost', 'field': k, 'role': 'planning_predictor', 'unit': 'dimensionless_score' if ('Cost' in k or k in ('Initial_Budget', 'Expected_Revenue', 'Contingency_Reserve')) else 'percentage_points' if 'Margin' in k else 'm2' if '_m2' in k else 'ratio' if 'Utilization' in k else 'hours' if 'Hours' in k else 'binary' if k.startswith('Event_Type_') else 'persons' if k in ('Guest_Count', 'Staff_Count') else 'seats' if k == 'Seating_Count' else 'person_hours', 'origin': 'synthetic_planning'} for k in FEATURES]
    dictionary += [{'dataset': 'XGBoost and graph', 'field': k, 'role': 'target_only', 'unit': 'binary', 'origin': 'simulated_outcome'} for k in TARGETS]
    dictionary += [{'dataset': 'Prophet', 'field': k, 'role': role, 'unit': unit, 'origin': 'publisher_observation_or_documented_aggregate'} for k, role, unit in [('ds', 'period_key', 'month_start'), ('y', 'measurement_target', 'see series unit; do not pool'), ('series_id', 'separate_model_identity', 'identifier'), ('source_id', 'provenance', 'identifier'), ('currency', 'unit_metadata', 'LKR or missing for index')]]
    write_csv(stage / 'reports/model_feature_dictionary.csv', dictionary)
    graphdir = stage / 'prepared/spatial_financial_graphs'
    graphdir.mkdir(parents=True)
    ledger = defaultdict(list)
    for r in bundle['line_items']:
        ledger[r['scenario_id']].append(r)
    source_by_id = {r['scenario_id']: r for r in provenance}
    y_by_id = {r['Scenario_ID']: r for r in xgb}
    nodes, edges, graphs, labels = [], [], [], []
    x = np.full((len(xgb), len(NODE_TYPES), len(NODE_FEATURES)), np.nan, dtype=np.float64)
    for graph_index, p in enumerate(bundle['planning']):
        sid = p['scenario_id']
        source = source_by_id[sid]
        values = graph_nodes(p, bundle['configuration']['financial_model'], ledger[sid])
        graphs.append({'graph_id': sid, 'Base_Event_ID': p['family_id'], 'split_group_id': source['split_group_id'], 'split': source['split'],
                       'node_count': len(NODE_TYPES), 'edge_count': len(EDGE_RULES), 'layout_origin': 'synthetic_assumption',
                       'Record_Kind': 'synthetic_event_financial', 'schema_version': 'unitless-dependency-graph.1.0',
                       'configuration_sha256': EXPECTED_HASH, 'planning_fingerprint': source['planning_fingerprint']})
        labels.append({'graph_id': sid, **{k: y_by_id[sid][k] for k in TARGETS}, 'label_origin': 'SIMULATOR_OUTCOME_NOT_OBSERVED'})
        for node_index, features in enumerate(values):
            nodes.append({'graph_id': sid, 'node_index': node_index, 'node_type': NODE_TYPES[node_index], **features})
            for column, name in enumerate(NODE_FEATURES):
                if features[name] is not None:
                    x[graph_index, node_index, column] = features[name]
        for edge_index, (start, end, relation, assumptions, parameter, unit) in enumerate(EDGE_RULES):
            edges.append({'graph_id': sid, 'edge_id': edge_index, 'source_node_index': start, 'target_node_index': end,
                          'relation_type': relation, 'rule_id': 'UNIT-' + relation, 'assumption_ids': assumptions,
                          'edge_origin': 'assumed_dependency', 'parameter': parameter, 'parameter_unit': unit})
    write_csv(graphdir / 'graph_index.csv', graphs)
    write_csv(graphdir / 'node_features.csv', nodes)
    write_csv(graphdir / 'edges.csv', edges)
    write_csv(graphdir / 'graph_targets.csv', labels)
    write_csv(graphdir / 'graph_provenance.csv', provenance)
    y = np.array([[r[k] for k in TARGETS] for r in labels], dtype=np.int64)
    edge_index = np.array([[r[0] for r in EDGE_RULES], [r[1] for r in EDGE_RULES]], dtype=np.int64)
    np.savez_compressed(graphdir / 'graphs.npz', x=x, feature_mask=np.isfinite(x), y=y, label_mask=np.ones_like(y, dtype=bool),
                        edge_index=edge_index, node_type_index=np.arange(len(NODE_TYPES)),
                        graph_id=np.array([r['graph_id'] for r in graphs]), split=np.array([r['split'] for r in graphs]),
                        split_group_id=np.array([r['split_group_id'] for r in graphs]))
    units = {k: 'm2' if k.endswith('_m2') else 'dimensionless_score' if k.endswith('_score') else 'hours' if 'hours' in k and k != 'weighted_person_hours' else 'overtime_weighted_person_hours' if k == 'weighted_person_hours' else 'metres' if k.endswith('_m') else 'percentage_points' if k == 'minimum_margin_points' else 'binary' if k in ('wedding', 'corporate') else 'persons' if k in ('guests', 'capacity_persons', 'staff') else k for k in NODE_FEATURES}
    graphschema = {'version': 'unitless-dependency-graph.1.0', 'adapted_from': 'reports/spatial_financial_graph_schema.md',
                   'node_types': NODE_TYPES, 'feature_names': NODE_FEATURES, 'feature_units': units, 'targets': TARGETS,
                   'serialization': 'Batched NPZ plus flat CSVs indexed by graph_id; no object/pickle arrays. edge_index 2x13 is shared by each graph.',
                   'missing_values': 'CSV empty / NPZ NaN with explicit feature_mask. Append node type one-hot using node_type_index before GNN use. Masked zeros in a later tensor loader are computational placeholders, not observed zeros.',
                   'disabled_node_scopes': 'VIP, booths, lighting, security and transport are out of this unitless scope; no fabricated nodes or measurements.',
                   'topology': 'One fixed functional topology, no coordinates, self-loops or invented reverse edges.',
                   'edge_rules': EDGE_RULES, 'source_units_amendment': 'Original proposed LKR wording replaced with dimensionless_score for this already approved unitless representation. Financial rules and records are unchanged.'}
    (graphdir / 'schema.json').write_text(json.dumps(graphschema, indent=2))
    # Read every newly written analytical table back and compare exact serialization.
    verification = {}
    for path, expected in [('prepared/real/weekly_retail_observations.csv', prices), ('prepared/real/monthly_inflation_observations.csv', indices),
                           ('prepared/real/observation_provenance.csv', links), ('prepared/prophet/prophet_cost_timeseries.csv', prophet),
                           ('prepared/prophet/chronological_split_manifest.csv', partitions), ('prepared/prophet/series_coverage.csv', coverage),
                           ('prepared/benchmarks/published_lkr_rate_evidence.csv', rates), ('prepared/benchmarks/provider_terms.csv', terms),
                           ('prepared/benchmarks/sl_procurement_context.csv', procurement), ('prepared/xgboost/xgboost_risk_dataset.csv', xgb),
                           ('prepared/xgboost/grouped_split_manifest.csv', splitrows), ('prepared/spatial_financial_graphs/graph_index.csv', graphs),
                           ('prepared/spatial_financial_graphs/node_features.csv', nodes), ('prepared/spatial_financial_graphs/edges.csv', edges),
                           ('prepared/spatial_financial_graphs/graph_targets.csv', labels), ('prepared/spatial_financial_graphs/graph_provenance.csv', provenance),
                           ('reports/model_feature_dictionary.csv', dictionary)]:
        encoded = [{k: '' if v is None else json.dumps(v, sort_keys=True) if isinstance(v, (list, dict)) else str(v) for k, v in r.items()} for r in expected]
        verification[path] = read_csv(stage / path) == encoded
    for split in ('train', 'validation', 'test'):
        verification['xgboost_split_' + split] = read_csv(stage / f'prepared/xgboost/{split}.csv') == [r for r in read_csv(stage / 'prepared/xgboost/xgboost_risk_dataset.csv') if r['split'] == split]
        verification['prophet_split_' + split] = {r['observation_id'] for r in read_csv(stage / f'prepared/prophet/{split}.csv')} == {r['observation_id'] for r in partitions if r['partition'] == split}
    with np.load(graphdir / 'graphs.npz', allow_pickle=False) as saved:
        verification['graph_npz_values'] = np.array_equal(saved['x'], x, equal_nan=True) and np.array_equal(saved['y'], y) and np.array_equal(saved['feature_mask'], np.isfinite(x)) and np.array_equal(saved['edge_index'], edge_index)
        verification['graph_split_identity'] = saved['split'].tolist() == [r['split'] for r in provenance]
    verification['graph_node_edge_scope'] = len(nodes) == len(xgb) * 9 and len(edges) == len(xgb) * 13 and all(0 <= r['source_node_index'] < 9 and 0 <= r['target_node_index'] < 9 for r in edges)
    verification['no_prediction_outcome_features'] = not FORBIDDEN.intersection(FEATURES) and not set(TARGETS).intersection(FEATURES)
    verification['synthetic_originals_unchanged'] = all(sha(ROOT / p) == value for p, value in protected.items())
    if not all(verification.values()):
        raise ValueError('output_verification_failed:' + str([k for k, v in verification.items() if not v]))
    # Preserve complete upstream evidence for portable reproduction; these snapshots are not extra selected model rows.
    snapshots = set(protected) | {str(p.relative_to(ROOT)).replace('\\', '/') for p in HISTORICAL.iterdir() if p.is_file()}
    snapshots |= set(hmanifest['source_hashes']) | set(hmanifest['publisher_hashes'])
    snapshots |= {str(pdf.relative_to(ROOT)).replace('\\', '/'), 'data/processed/sl_procurement.csv', 'data/raw/extractions/SL_Procurement.csv',
                  'reports/public_lkr_rate_evidence.xlsx', 'reports/parameter_evidence_summary.md', 'reports/spatial_financial_graph_schema.md',
                  'reports/financial_dependency_rules.md', 'reports/unitless_formula_substantive_review.md', 'reports/updated_data_dictionary.csv',
                  'reports/simulation_assumption_register.csv', 'config/unitless_bulk_candidate.v0.1.proposed.json',
                  'config/unitless_bulk_authorization.v0.1.user.json', 'reports/unitless_bulk_authorization_v0.1/user_bulk_authorization_evidence.json',
                  'reports/unitless_bulk_execution_20261012/validation_results.json', 'reports/unitless_bulk_execution_20261012/validation_report.md'}
    snapshots |= {str(p.relative_to(ROOT)).replace('\\', '/') for parent in ('component4', 'tests') for p in (ROOT / parent).glob('*.py')}
    snapshots |= {'scripts/research_pipeline/create_hybrid_dataset.py', 'scripts/research_pipeline/features.py', 'scripts/research_pipeline/common.py',
                  'scripts/research_pipeline/execute_unitless_bulk.py', 'work/minimal_pilot_authoring/export_unitless_bulk.mjs',
                  'work/minimal_pilot_authoring/export_unitless.mjs'}
    for name in sorted(snapshots):
        destination = stage / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(ROOT / name, destination)
        if sha(destination) != sha(ROOT / name):
            raise ValueError('snapshot_copy_changed:' + name)
    statistics = {'synthetic_event_scenarios': len(xgb), 'synthetic_root_families': len({r['family_id'] for r in provenance}),
                  'genuine_historical_measurements_selected': len(raw_ids), 'direct_weekly_prices': len(prices), 'direct_monthly_indices': len(indices),
                  'published_advertisement_items_not_transactions': len(rates), 'aggregate_procurement_context_records': len(procurement),
                  'real_observed_event_financial_outcomes': 0, 'prophet_ready_records': len(prophet), 'prophet_independent_series': len(coverage),
                  'prophet_monthly_price_aggregates': sum(r['source_id'] == 'S003' for r in prophet), 'xgboost_ready_rows': len(xgb),
                  'xgboost_predictor_count': len(FEATURES), 'graph_count': len(graphs), 'node_records': len(nodes), 'edge_records': len(edges),
                  'distinct_planning_fingerprints': len({r['planning_fingerprint'] for r in provenance}), 'fixed_graph_topologies': 1,
                  'split_rows': dict(Counter(r['split'] for r in provenance)),
                  'split_families': dict(Counter({r['family_id']: r['split'] for r in provenance}.values())),
                  'split_groups': dict(Counter({r['split_group_id']: r['split'] for r in provenance}.values())),
                  'prophet_partitions': dict(Counter(r['partition'] for r in partitions)),
                  'historical_quality_review_measurements': sum(bool(r['quality_flag']) for r in prices + indices),
                  'prophet_quality_review_records': sum(r['source_quality_review'] == '1' for r in prophet),
                  'target_class_counts': {k: dict(Counter(str(r[k]) for r in xgb)) for k in TARGETS}}
    (stage / 'reports/test_results.json').write_text(json.dumps(test_receipt, indent=2))
    (stage / 'reports/validation_results.json').write_text(json.dumps({'status': 'PASS_FOR_INTERNAL_RESEARCH_DATA_PREPARATION',
        'checks': verification, 'statistics': statistics, 'synthetic_original_sha256': protected, 'group_split_audit': split_audit,
        'runtime': {'python': platform.python_version(), 'numpy': np.__version__}, 'model_training_performed': False,
        'real_world_financial_validation': False, 'qualified_formula_reviews_pending': 19}, indent=2))
    (stage / 'README.md').write_text(documentation(statistics), encoding='utf-8')
    (stage / 'reports/data_leakage_audit.md').write_text(leakage_document(), encoding='utf-8')
    (stage / 'requirements.txt').write_text('numpy==2.3.5\nopenpyxl==3.1.5\n', encoding='utf-8')
    payload_hashes = {str(p.relative_to(stage)).replace('\\', '/'): sha(p) for p in sorted(stage.rglob('*')) if p.is_file()}
    (stage / 'checksums.sha256').write_text(''.join(f'{value}  {name}\n' for name, value in payload_hashes.items()), encoding='utf-8')
    (stage / 'package_manifest.json').write_text(json.dumps({'version': 'Component4_Hybrid_Dataset_Final_v1.0',
        'status': 'VALIDATED_INTERNAL_EXPERIMENTAL_HYBRID_PACKAGE_NOT_REAL_WORLD_EVENT_VALIDATION',
        'frozen_synthetic_configuration_sha256': EXPECTED_HASH, 'statistics': statistics, 'files': payload_hashes,
        'snapshot_role': 'Upstream full source tables are reproduction evidence, not additional selected model observations.',
        'new_synthetic_scenarios_generated': 0, 'model_training_performed': False}, indent=2))
    if any(sha(ROOT / p) != value for p, value in protected.items()):
        raise ValueError('protected_synthetic_directory_changed_before_commit')
    return statistics


def documentation(s):
    return f'''# Component 4 hybrid research dataset v1.0

Prepared from existing datasets without regenerating or modifying a single synthetic scenario. No models or frontend were trained or rebuilt.

## Contents and counts

- Original synthetic master CSV/XLSX and complete provenance/ledger/configuration: `data/final/unitless-bulk-v0.1/`. {s['synthetic_event_scenarios']} scenarios, 1250 roots, 3750 planning states and 2500 paired outcome streams. Scores are dimensionless, never LKR.
- Selected real measurements: `prepared/real/`. {s['direct_weekly_prices']} weekly retail measurements and {s['direct_monthly_indices']} monthly inflation indices, {s['genuine_historical_measurements_selected']} direct historical measurements in total.
- Prophet views: `prepared/prophet/`. {s['prophet_ready_records']} records in {s['prophet_independent_series']} distinct monthly series: 450 documented weekly-price aggregates plus 113 index observations. These reuse the real measurements; do not add them to an independent observation count.
- Advertised evidence: `prepared/benchmarks/`. 20 previously recorded Sri Lankan tariff items, nine provider terms, and 60 aggregate procurement context rows. These are not observed event-financial outcomes, complete quotations or calibrated simulator inputs.
- XGBoost: `prepared/xgboost/`. {s['xgboost_ready_rows']} rows, 21 planning-time predictors, three simulated targets, exact original grouped assignments. IDs/splits/hashes are metadata, not predictors. No fitted transformations.
- Graphs: `prepared/spatial_financial_graphs/`. {s['graph_count']} graphs, {s['node_records']} node records and {s['edge_records']} directed edges. Batched NPZ and flat CSVs are an explicit serialization adaptation of the existing proposed graph schema. Arrays contain masks and node-type indices. The nine-node, thirteen-edge topology is a designed rule dependency network, not measured adjacency or CAD.
- Real observed event-financial outcomes: **zero**.

The source snapshot tables and full PDFs contain additional unselected rows for provenance and reproducibility. Only the `prepared/` tables define the selected real/model-ready datasets. Do not concatenate source snapshots, monthly aggregates and synthetic master rows into a supposedly homogeneous event dataset.

## Selection, dates and units

Ten full catering-relevant ingredient histories were selected before any model training: rice, lentils, wheat flour, sugar, coconut oil, broiler poultry, eggs, onions, potatoes and bananas. Their exact publisher names/units are in `component4/hybrid_dataset.py:BASKET`. Both CCPI and NCPI retain their own base-2021 series. No arbitrary row truncation, imputation or unit conversion was used.

Weekly observations retain original week labels, source record IDs, extraction rows/hashes and publisher URLs/hashes. Exact weekly observation/publication dates are not known and are not fabricated. Monthly `ds` means the first day of a measurement period, not a release timestamp. Price aggregates are means of the source weeks actually collected; complete publisher coverage and publication vintages are not established. The open October 2026 month is excluded. Source quality flags are preserved ({s['historical_quality_review_measurements']} direct records; {s['prophet_quality_review_records']} monthly records need review). Arithmetic, lineage and byte checks do not independently establish perfect PDF transcription.

Do not pool index points and LKR ingredient prices into a single Prophet target. The 70/15/15 chronological splits are inherited, per-series univariate retrospective assignments, with at least 24 training and six validation/test months. Different series have different cutoffs; no pooled temporal model is authorized. Fit any later data-dependent transformations on training data only.

The procurement file has overlapping reporting periods and mixed average/period-total semantics. It is context only and cannot produce per-event or per-service prices. Foreign California venues and UK/US prices are excluded from local modeling. Advertised LKR rates retain pending reviews, unknown tax/effective dates, exclusions and bundled-service restrictions. No new primary-page archival verification was performed; evidence is the retained access-date-2026-10-11 register, not a current quotation.

## Splits and leakage

Synthetic train/validation/test rows: 4500/1500/1500; root families: 750/250/250; connected split groups: 711/248/248. Original split assignments were not changed. Graphs inherit these assignments exactly. Repeated outcome replications share planning features and remain in one group. Shared venue/vendor/source assumptions are intentionally disclosed; unseen-template evaluation is unavailable. The {s['distinct_planning_fingerprints']} distinct planning fingerprints are not independent real events.

Actual costs/revenue/profit/margins, actual budget utilization, actual cost per guest/variance, latent shocks and anomaly injection state are excluded from prediction inputs. Planned rule-based cost is allowed because it was computed before shocks, not fitted to future outcomes. Historical prices are **not joined** to synthetic score scenarios or converted into LKR calibration. See `reports/data_leakage_audit.md` and feature/schema specifications.

Vendor anomaly is a random simulation injection independent of planning inputs. Its label can support pipeline/negative-control experiments, but this feature set has no substantive signal for predicting that random injection. Synthetic performance cannot validate real-world risk prediction. All 19 formulas still lack independent qualified review. The existing bulk engineering waiver is preserved, not expanded into professional or institutional approval.

## Verification and reproduction

All original synthetic file checksums are recorded in `reports/validation_results.json`; the original CSV/XLSX copies are byte-identical. `checksums.sha256` and `package_manifest.json` verify packaged content. Python 3.12.14, NumPy 2.3.5 and read-only openpyxl 3.1.5 prepared this package. The full project test run is recorded in `reports/test_results.json`.

From the ORIGINAL project in Windows PowerShell, using a fresh destination:

```powershell
$taskPython='C:/Users/User/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/python.exe'
& $taskPython scripts/research_pipeline/create_hybrid_dataset.py --destination data/hybrid/hybrid_reproduction.v1.0 --archive outputs/Component4_Hybrid_Dataset_Final_v1.0_replay.zip
```

This only reads existing scenarios and prepares derivatives. It never invokes a scenario generator. The archive includes upstream evidence and preparation source. Full legacy test execution requires the original project environment and auxiliary fixtures; the portable package verification below requires only the two requirements-listed dependencies.

After extracting the archive:

```powershell
python scripts/research_pipeline/verify_hybrid_package.py .
```

Neither ZIP byte equality nor NPZ container timestamps imply dataset reproducibility. Compare CSV values and source hashes. Any later model experiment must preserve these fixed groups and report its experimental/synthetic limitations. No model training was performed.
'''


def leakage_document():
    return '''# Hybrid dataset leakage audit

PASS for preparation-stage checks under the documented conditional simulator scope.

XGBoost uses a fixed 21-field planning allowlist. The master Actual_Total_Cost, Actual_Revenue, Budget_Utilization, Profit, Profit_Margin, Cost_Per_Guest, Cost_Variance, Financial_Health_Score and all hidden shocks are excluded. Targets are stored separately from the input-feature specification; select only its predictors when training. IDs, configuration hashes, family/split fields, planning state and outcome-replication metadata are not predictors. Forecasted_Total_Cost is the existing pre-shock deterministic estimate, not a fitted Prophet or ML output.

Original family/fingerprint connected components and 60/20/20 assignments are preserved. Matching planning fingerprints and all paired descendants remain within a single split. Shared source/venue/vendor ancestry crosses splits by design; conditional in-simulator evaluation is the only supported claim. Strict unseen-template independence is not satisfied and is not claimed. Related rows do not create 7500 independent samples.

Prophet series retain separate commodity/unit/index/source identities. Training dates precede validation and test dates within each series, with unique month keys and no internal missing months. Splits precede any fitted transformations. The inputs are retrospective monthly aggregates; observation-release timestamps and vintages remain unknown. Do not use a whole-month average for an event earlier in that month. There is no historical-to-synthetic join in this package. Real-time forecasting or pooled cross-series cutoffs require a later design.

Graph inputs contain only planned quantities, area constraints, ex-ante financial context and disjoint planned service charges. Labels and simulated outcomes are not node/edge features. NPZ masks distinguish absent features from valid zeros. Graph IDs and target tables join exactly to the master and inherited split manifests. One fixed assumed topology does not establish learned physical adjacency, causal structure or unseen-venue performance. Original formula versions and cost units remain unchanged.

Random vendor anomaly injections are independent of the permissible feature set. Do not leak the injection state to obtain artificial performance. No scaling, imputation, SMOTE, feature selection, target balancing, hyperparameter search or model fitting was performed. Model-specific cross-validation must group within training and preserve the untouched test partition.
'''


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--destination', required=True)
    parser.add_argument('--archive', required=True)
    args = parser.parse_args()
    destination = (ROOT / args.destination).resolve()
    archive = (ROOT / args.archive).resolve()
    if not destination.is_relative_to(ROOT / 'data/hybrid') or destination == (ROOT / 'data/hybrid') or not archive.is_relative_to(ROOT / 'outputs'):
        raise ValueError('new_scoped_paths_required')
    if destination.exists() or archive.exists():
        raise ValueError('refuse_to_overwrite_existing_output')
    environment = dict(os.environ, TEMP=str(ROOT / 'work'), TMP=str(ROOT / 'work'))
    tests = subprocess.run([sys.executable, '-m', 'unittest', 'discover', '-s', 'tests', '-q'], cwd=ROOT, env=environment, text=True, capture_output=True)
    receipt = {'command': 'python -m unittest discover -s tests -q', 'exit_code': tests.returncode,
               'stdout': tests.stdout, 'stderr': tests.stderr, 'date_local': '2026-10-12', 'models_trained': False}
    if tests.returncode:
        print(tests.stdout + tests.stderr)
        raise SystemExit('Tests failed; no hybrid package published')
    destination.parent.mkdir(parents=True, exist_ok=True)
    archive.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.hybrid-staging-', dir=destination.parent) as temporary:
        stage = Path(temporary) / 'package'
        stage.mkdir()
        stats = prepare(stage, receipt)
        # The independent portable checker must pass before the atomic publish.
        checker = ROOT / 'scripts/research_pipeline/verify_hybrid_package.py'
        shutil.copy2(checker, stage / 'scripts/research_pipeline/verify_hybrid_package.py')
        # Register the checker in the package checksums after copying it.
        mpath = stage / 'package_manifest.json'
        m = json.loads(mpath.read_text())
        m['files']['scripts/research_pipeline/verify_hybrid_package.py'] = sha(checker)
        mpath.write_text(json.dumps(m, indent=2))
        (stage / 'checksums.sha256').write_text(''.join(f'{value}  {name}\n' for name, value in sorted(m['files'].items())))
        subprocess.run([sys.executable, str(checker), str(stage)], check=True, cwd=ROOT)
        os.rename(stage, destination)
    temporary_zip = archive.with_suffix('.zip.part')
    if temporary_zip.exists():
        raise ValueError('existing_archive_staging_file')
    with zipfile.ZipFile(temporary_zip, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=1) as package:
        for path in sorted(destination.rglob('*')):
            if path.is_file():
                package.write(path, str(path.relative_to(destination)).replace('\\', '/'))
    with zipfile.ZipFile(temporary_zip) as package:
        if package.testzip() is not None:
            raise ValueError('zip_crc_failure')
        expected = json.loads((destination / 'package_manifest.json').read_text())['files']
        import hashlib
        for name, digest in expected.items():
            if hashlib.sha256(package.read(name)).hexdigest() != digest:
                raise ValueError('archive_entry_checksum:' + name)
    os.rename(temporary_zip, archive)
    receipt_path = archive.with_suffix('.zip.sha256')
    receipt_path.write_text(sha(archive) + '  ' + archive.name + '\n')
    print(json.dumps({'status': 'PASS', 'package': str(destination), 'archive': str(archive),
                      'archive_sha256': sha(archive), 'statistics': stats, 'tests': tests.stderr.strip()}, indent=2))


if __name__ == '__main__':
    main()

"""Mock fixtures only: no authorization is granted and no project pilot is written."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from component4.unitless_pilot import (ROOT, validate, generate, generate_test,
                                      execution_blockers, validate_bundle, _write_atomic)
from component4.engine import fingerprint

NODE = Path('C:/Users/User/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe')


def candidate(mock=False):
    c = json.loads((ROOT / 'config/unitless_pilot_candidate.json').read_text())
    if mock:
        c['record_kind'] = 'mock_unitless_test'
        c['origin'] = 'ISOLATED_AUTOMATED_TEST_FIXTURE_NOT_APPROVED_PILOT'
    return c


class UnitlessPipelineTests(unittest.TestCase):
    def test_candidate_is_complete_but_unapproved(self):
        self.assertTrue(validate(candidate()))
        self.assertIn('separate_explicit_authorization_missing', execution_blockers(candidate()))

    def test_no_authorization_no_generation(self):
        with self.assertRaisesRegex(ValueError, 'execution_blocked'):
            generate(candidate(), {})

    def test_incorrect_hash(self):
        self.assertIn('exact_configuration_hash_not_authorized', execution_blockers(candidate(), {'configuration_sha256': '0'*64}))

    def test_mock_cannot_use_authorized_path(self):
        self.assertIn('test_mock_forbidden_in_authorized_execution', execution_blockers(candidate(True)))

    def test_live_candidate_cannot_activate_test_path(self):
        with self.assertRaisesRegex(ValueError, 'explicit_test_only'):
            generate_test(candidate())

    def test_missing_parameter_blocks(self):
        c = candidate(); c['coefficients']['catering'] = None
        with self.assertRaises(ValueError): validate(c)

    def test_bool_integer_rejected(self):
        c = candidate(); c['capacity_persons'] = True
        with self.assertRaises(ValueError): validate(c)

    def test_currency_conversion_rejected(self):
        c = candidate(); c['currency'] = 'LKR'
        with self.assertRaises(ValueError): validate(c)

    def test_invalid_quotas(self):
        c = candidate(); c['proposed_pilot']['archetype_root_quotas']['Wedding'] = 4
        with self.assertRaises(ValueError): validate(c)

    def test_source_snapshot_tampering(self):
        c = candidate(); c['experimental_source_snapshot']['version'] = 'changed'
        with self.assertRaisesRegex(ValueError, 'snapshot_mismatch'): validate(c)

    def test_reproducibility_and_ids(self):
        a = generate_test(candidate(True)); b = generate_test(candidate(True))
        self.assertEqual(a,b)
        self.assertEqual(len({r['scenario_id'] for r in a['planning']}), 30)
        self.assertEqual(len({r['family_id'] for r in a['planning']}), 5)

    def test_canonical_key_order_preserves_ids_and_planning(self):
        c = candidate(True)
        reordered = deepcopy(c)
        reordered['proposed_pilot']['archetype_root_quotas'] = {'Corporate': 2, 'Wedding': 3}
        reordered['profiles'] = dict(reversed(list(reordered['profiles'].items())))
        reordered['coefficients'] = dict(reversed(list(reordered['coefficients'].items())))
        reordered = dict(reversed(list(reordered.items())))
        self.assertEqual(fingerprint(c), fingerprint(reordered))
        self.assertEqual(generate_test(c), generate_test(reordered))

    def test_arithmetic_and_synthetic_labels(self):
        b = generate_test(candidate(True)); self.assertTrue(validate_bundle(b))
        self.assertTrue(all(x['origin'].startswith('MOCK_') for x in b['provenance']))
        self.assertTrue(all(x['label_origin'] == 'SIMULATOR_OUTCOME_NOT_OBSERVED' for x in b['outcomes']))

    def test_no_latent_feature_leakage(self):
        b = generate_test(candidate(True))
        forbidden = {'simulated_cost_score', 'simulated_profit_score', 'cost_multiplier', 'receipt_multiplier', 'injection_multiplier', 'simulated_budget_overrun', 'simulated_vendor_cost_anomaly'}
        self.assertFalse(forbidden & set(b['planning'][0]))

    def test_paired_common_random_numbers(self):
        b = generate_test(candidate(True)); lookup = {x['scenario_id']: x for x in b['latent_DO_NOT_USE_AS_FEATURES']}
        for family in {x['family_id'] for x in b['planning']}:
            for rep in (0,1):
                shocks = [{k:v for k,v in lookup[x['scenario_id']].items() if k!='scenario_id'} for x in b['planning'] if x['family_id']==family and x['outcome_replication']==rep]
                self.assertEqual(shocks[0],shocks[1]); self.assertEqual(shocks[0],shocks[2])

    def test_group_safe_split_shared_venue_retained(self):
        b = generate_test(candidate(True))
        self.assertEqual(b['manifest']['split_audit']['status'], 'PASS')
        self.assertEqual(len({x['venue_template_id'] for x in b['provenance']}),1)
        self.assertFalse(b['manifest']['split_audit']['unseen_venue_claim_supported_by_protocol'])
        for f in {x['family_id'] for x in b['provenance']}:
            self.assertEqual(len({x['split'] for x in b['provenance'] if x['family_id']==f}),1)

    def test_arithmetic_corruption_rejected(self):
        b=generate_test(candidate(True)); b['outcomes'][0]['simulated_cost_score']+=1
        with self.assertRaisesRegex(ValueError,'reconciliation'): validate_bundle(b)

    def test_double_charging_rejected(self):
        b=generate_test(candidate(True)); b['line_items'].append(deepcopy(b['line_items'][0]))
        with self.assertRaisesRegex(ValueError,'double_charging'): validate_bundle(b)

    def test_geometry_corruption_rejected(self):
        b=generate_test(candidate(True)); b['planning'][0]['used_area_m2']+=1
        with self.assertRaisesRegex(ValueError,'geometry'): validate_bundle(b)

    def test_label_corruption_rejected(self):
        b=generate_test(candidate(True)); b['outcomes'][0]['simulated_budget_overrun']=not b['outcomes'][0]['simulated_budget_overrun']
        with self.assertRaisesRegex(ValueError,'label'): validate_bundle(b)

    def test_failed_export_is_atomic(self):
        b=generate_test(candidate(True))
        def fail(*args): raise RuntimeError('test injected export failure')
        with tempfile.TemporaryDirectory(dir=ROOT / 'work') as d:
            target=Path(d)/'mock-output'
            with self.assertRaises(RuntimeError): _write_atomic(b,target,NODE,exporter=fail)
            self.assertFalse(target.exists()); self.assertEqual(list(Path(d).iterdir()),[])

    def test_existing_destination_not_overwritten(self):
        with tempfile.TemporaryDirectory(dir=ROOT / 'work') as d:
            with self.assertRaisesRegex(ValueError,'new_destination'): _write_atomic(generate_test(candidate(True)),d,NODE)

    def test_real_csv_xlsx_export_in_temporary_mock_directory(self):
        from openpyxl import load_workbook  # Read-only independent export QA, not authoring.
        with tempfile.TemporaryDirectory(dir=ROOT / 'work') as d:
            target=_write_atomic(generate_test(candidate(True)),Path(d)/'mock-only',NODE)
            workbook=load_workbook(target/'unitless_dataset.xlsx',read_only=True)
            try:
                self.assertEqual(sum(1 for _ in workbook['planning'].iter_rows()),31)
                self.assertEqual(sum(1 for _ in workbook['line_items'].iter_rows()),211)
            finally:
                workbook.close()
            manifest=json.loads((target/'manifest.json').read_text())
            for name, sha in manifest['files'].items(): self.assertEqual(hashlib.sha256((target/name).read_bytes()).hexdigest(),sha)
            self.assertTrue((target/'provenance.csv').exists())


if __name__ == '__main__': unittest.main()
